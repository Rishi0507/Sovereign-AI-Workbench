"""End-to-end API tests with the real approval queue and job scheduler (PRD phase 9)."""

from __future__ import annotations

import re
import time
from collections.abc import Callable, Iterator
from typing import Any

import pytest
from fastapi.testclient import TestClient

from workbench.agent.approvals import QueueApprove
from workbench.api.app import create_app
from workbench.runtime import Runtime

ENG = {"X-User": "engineer1"}
OFFICER = {"X-User": "officer1"}
OWNER = {"X-User": "owner1"}


@pytest.fixture()
def api(make_runtime: Callable[..., Runtime]) -> Iterator[tuple[TestClient, Runtime]]:
    # Every gate is exercised explicitly here; the shortcuts have their own tests below.
    rt = make_runtime(approvals=QueueApprove(poll_s=0.05), plan_approval_covers_drafts=False)
    rt.jobs.start()
    app = create_app(rt, install_guard=False)
    with TestClient(app) as client:
        yield client, rt


@pytest.fixture()
def api_default(make_runtime: Callable[..., Runtime]) -> Iterator[tuple[TestClient, Runtime]]:
    rt = make_runtime(approvals=QueueApprove(poll_s=0.05))
    rt.jobs.start()
    app = create_app(rt, install_guard=False)
    with TestClient(app) as client:
        yield client, rt


def wait_for(client: TestClient, task_id: str, cond: Callable[[dict[str, Any]], bool], timeout: float = 60,
              headers: dict[str, str] = ENG) -> dict[str, Any]:
    deadline = time.time() + timeout
    while time.time() < deadline:
        task = client.get(f"/api/tasks/{task_id}", headers=headers).json()
        if cond(task):
            return task
        time.sleep(0.05)
    raise AssertionError(f"timed out; last status {task['status']} gate {task.get('pending_gate')}")


def gate_is(kind: str) -> Callable[[dict[str, Any]], bool]:
    return lambda t: bool(t["pending_gate"]) and t["pending_gate"]["kind"] == kind


def run_to_draft(client: TestClient, text: str, attachments: list[str]) -> dict[str, Any]:
    resp = client.post("/api/tasks", headers=ENG, json={"workspace": "plant-a", "text": text,
                                                         "attachments": attachments})
    assert resp.status_code == 201, resp.text
    task_id = resp.json()["id"]
    task = wait_for(client, task_id, gate_is("plan"))
    assert task["plan"]["template"] == "approval_note_from_scan"
    assert client.post(f"/api/tasks/{task_id}/plan/decision", headers=ENG, json={"decision": "approve"}).status_code == 200
    task = wait_for(client, task_id, gate_is("action"))
    gate = task["pending_gate"]
    assert gate["payload"]["tool"] == "make_docx"
    r = client.post(f"/api/tasks/{task_id}/actions/{gate['id']}/decision", headers=ENG, json={"approve": True})
    assert r.status_code == 200
    return wait_for(client, task_id, gate_is("deliverable"))


def test_trace_a_through_the_api(api: tuple[TestClient, Runtime]) -> None:
    client, rt = api
    task = run_to_draft(client, "Draft an approval note for this inspection report", ["inputs/inspection_P108B.pdf"])
    task_id = task["id"]
    assert task["label_display"] == "Confidential"
    assert task["route"]["chosen"] == "qwen3-vl-8b"
    mismatches = [c for c in task["checks"] if c["status"] == "mismatch"]
    assert sorted(c["rule"] for c in mismatches) == ["row_tags_match", "thickness_trend", "thickness_vs_limit"]

    draft = client.get(f"/api/tasks/{task_id}/draft", headers=ENG).json()
    assert draft["blockers"] == ["3 mismatch(es) not acknowledged"]
    blocks = draft["deliverables"][0]["preview"]["blocks"]
    assert blocks[0]["type"] == "marking" and blocks[0]["text"] == "CONFIDENTIAL"
    assert draft["deliverables"][0]["provenance"]["counts"]["unsourced"] == 0

    refused = client.post(f"/api/tasks/{task_id}/draft/decision", headers=ENG, json={"decision": "approve"})
    assert refused.status_code == 409 and "not acknowledged" in refused.json()["detail"]
    assert client.post(f"/api/tasks/{task_id}/draft/decision", headers=ENG,
                       json={"decision": "reject"}).status_code == 422

    checks = client.get(f"/api/tasks/{task_id}/checks", headers=ENG).json()
    tag = next(d for d in checks["dual_read"] if d["body"].get("field_kind") == "tag")
    assert tag["body"]["value"] == "P-108B" and tag["confidence"] == "medium"
    crop = client.get(f"/api/tasks/{task_id}/evidence/{tag['body']['crop']}", headers=ENG)
    assert crop.status_code == 200 and crop.headers["content-type"] == "image/png"
    assert client.get(f"/api/tasks/{task_id}/evidence/..%2Fsecret.png", headers=ENG).status_code == 404

    for check in mismatches[:2]:
        r = client.post(f"/api/tasks/{task_id}/checks/{check['record_id']}/acknowledge", headers=ENG)
        assert r.status_code == 200
    ok = client.post(f"/api/tasks/{task_id}/draft/decision", headers=ENG,
                     json={"decision": "approve", "acknowledged": [mismatches[2]["record_id"]]})
    assert ok.status_code == 200, ok.text
    done = wait_for(client, task_id, lambda t: t["status"] == "completed")
    final_id = done["deliverables"][0]["final_file_id"]
    meta = client.get(f"/api/files/{final_id}", headers=ENG).json()
    assert meta["path"] == f"final/{task_id}/approval-note.docx" and meta["label_display"] == "Confidential"
    download = client.get(f"/api/files/{final_id}/download", headers=ENG)
    assert download.status_code == 200 and download.content[:2] == b"PK"

    share = client.post(f"/api/files/{final_id}/share", headers=ENG, json={"workspace": "plant-a-general"})
    assert share.status_code == 403 and "ceiling" in share.json()["detail"]
    types = [e.event["type"] for e in rt.audit.entries()]
    assert "file.share_refused" in types and "deliverable.final" in types and "gate.plan" in types
    assert client.get("/api/audit/verify").json()["ok"]

    record = client.get(f"/api/records/{mismatches[0]['record_id']}", headers=ENG).json()
    assert record["kind"] == "check_result" and len(record["chain"]) >= 2
    routing = client.get(f"/api/routing/{task_id}", headers=ENG).json()
    assert "threshold=0.80" in routing["log_line"]
    ledger = client.get(f"/api/tasks/{task_id}/ledger", headers=ENG).json()
    assert ledger[0]["kind"] == "user_input" and ledger[0]["trust"] == "control"

    # downgrade: requester cannot self-approve; approval re-stamps the file and keeps the reason
    req = client.post(f"/api/files/{final_id}/downgrade", headers=OFFICER,
                      json={"level": "Restricted", "reason": "commercial figures removed"})
    assert req.status_code == 201, req.text
    rid = req.json()["id"]
    assert client.post(f"/api/files/{final_id}/downgrade", headers=ENG,
                       json={"level": "Restricted", "reason": "x"}).status_code == 403
    assert client.post(f"/api/downgrades/{rid}/decision", headers=OFFICER, json={"approve": True}).status_code == 403
    assert client.post(f"/api/downgrades/{rid}/decision", headers=ENG, json={"approve": True}).status_code == 403
    decided = client.post(f"/api/downgrades/{rid}/decision", headers=OWNER, json={"approve": True})
    assert decided.status_code == 200 and decided.json()["status"] == "approved"
    assert client.get(f"/api/files/{final_id}", headers=ENG).json()["label_display"] == "Restricted"
    from docx import Document

    doc = Document(str(rt.files.open_path(final_id)))
    assert doc.sections[0].header.paragraphs[0].text == "RESTRICTED"
    event = next(e.event for e in rt.audit.entries() if e.event["type"] == "downgrade.approved")
    assert event["before"] == "Confidential" and event["after"] == "Restricted"
    assert event["reason"] == "commercial figures removed" and event["before_sha256"] != event["after_sha256"]
    assert client.post(f"/api/downgrades/{rid}/decision", headers=OWNER, json={"approve": True}).status_code == 403
    shared = client.post(f"/api/files/{final_id}/share", headers=ENG, json={"workspace": "plant-a-general"})
    assert shared.status_code == 200


def test_orphan_figures_must_be_resolved(api: tuple[TestClient, Runtime]) -> None:
    client, rt = api
    task = run_to_draft(client, "Draft an approval note for this inspection report", ["inputs/inspection_P108B.pdf"])
    task_id = task["id"]
    from docx import Document

    state = rt.tasks.get(task_id)
    path = rt.files.open_path(state.deliverables[0].file_id)
    doc = Document(str(path))
    doc.add_paragraph("A further reading of 7.9 mm was reported by phone.")
    doc.add_paragraph("A second unverified reading of 8.4 mm was mentioned.")
    doc.save(str(path))
    rt.orchestrator.review(state)
    draft = client.get(f"/api/tasks/{task_id}/draft", headers=ENG).json()
    d = draft["deliverables"][0]
    orphans = [f for f in d["provenance"]["figures"] if f["status"] == "unsourced"]
    assert [f["raw"] for f in orphans] == ["7.9 mm", "8.4 mm"]
    assert "2 unsourced figure(s) not resolved" in draft["blockers"]
    key1, key2 = f"{d['file_id']}:{orphans[0]['id']}", f"{d['file_id']}:{orphans[1]['id']}"
    ocr = next(r for r in rt.ledger.for_task(task_id) if r.kind == "ocr_text")
    bad_link = client.post(f"/api/tasks/{task_id}/draft/figures/{key1}", headers=ENG,
                           json={"action": "link", "record_id": ocr.id})
    assert bad_link.status_code == 422
    fixed = client.post(f"/api/tasks/{task_id}/draft/figures/{key1}", headers=ENG,
                        json={"action": "correct", "value": "5.6 mm"})
    assert fixed.status_code == 200, fixed.text
    assert client.post(f"/api/tasks/{task_id}/draft/figures/{key2}", headers=ENG,
                       json={"action": "confirm", "note": "confirmed by the inspector"}).status_code == 200
    doc = Document(str(path))
    assert any("further reading of 5.6 mm" in p.text for p in doc.paragraphs)
    blockers = client.get(f"/api/tasks/{task_id}/draft", headers=ENG).json()["blockers"]
    assert blockers == ["3 mismatch(es) not acknowledged"]
    acks = [c["record_id"] for c in rt.tasks.get(task_id).checks if c["status"] == "mismatch"]
    assert client.post(f"/api/tasks/{task_id}/draft/decision", headers=ENG,
                       json={"decision": "approve", "acknowledged": acks}).status_code == 200
    wait_for(client, task_id, lambda t: t["status"] == "completed")
    assert any(e.event["type"] == "figure.decision" and e.event["action"] == "confirm" for e in rt.audit.entries())


def test_revision_request_reruns_the_draft(api: tuple[TestClient, Runtime]) -> None:
    client, rt = api
    task = run_to_draft(client, "Draft an approval note for this inspection report", ["inputs/inspection_P108B.pdf"])
    task_id = task["id"]
    r = client.post(f"/api/tasks/{task_id}/draft/decision", headers=ENG,
                    json={"decision": "reject", "note": "state the vendor explicitly"})
    assert r.status_code == 200
    task = wait_for(client, task_id, gate_is("action"))
    client.post(f"/api/tasks/{task_id}/actions/{task['pending_gate']['id']}/decision", headers=ENG,
                json={"approve": True})
    task = wait_for(client, task_id, gate_is("deliverable"))
    notes = [r for r in rt.ledger.for_task(task_id) if r.kind == "user_input"]
    assert any("state the vendor explicitly" in r.body_text() for r in notes)
    assert rt.tasks.get(task_id).revisions == 1


def test_edit_and_reject_plan(api: tuple[TestClient, Runtime]) -> None:
    client, _rt = api
    resp = client.post("/api/tasks", headers=ENG, json={"workspace": "plant-a", "text": "Summarise this vendor contract",
                                                         "attachments": ["inputs/vendor_contract.pdf"]})
    task_id = resp.json()["id"]
    task = wait_for(client, task_id, gate_is("plan"))
    typed = {"steps": [{"id": "read", "action": {"tool": "read_document"}, "args": {}, "inputs": [
        {"from": "attachment", "ref": "inputs/vendor_contract.pdf"}], "output_type": "document", "side_effect": False},
        {"id": "sheet", "action": {"tool": "make_xlsx"}, "args": {}, "inputs": [{"from": "step", "ref": "nothing"}],
         "output_type": "file", "side_effect": True}], "deliverables": [{"type": "xlsx"}]}
    assert client.post(f"/api/tasks/{task_id}/plan/decision", headers=ENG,
                       json={"decision": "edit", "plan": typed}).status_code == 200
    task = wait_for(client, task_id, lambda t: gate_is("plan")(t) and t["plan"]["errors"])
    assert any('input "nothing"' in e for e in task["plan"]["errors"])
    blocked = client.post(f"/api/tasks/{task_id}/plan/decision", headers=ENG, json={"decision": "approve"})
    assert blocked.status_code == 409
    assert client.post(f"/api/tasks/{task_id}/plan/decision", headers={"X-User": "buyer1"},
                       json={"decision": "reject"}).status_code == 403
    assert client.post(f"/api/tasks/{task_id}/plan/decision", headers=ENG,
                       json={"decision": "reject", "note": "not needed"}).status_code == 200
    wait_for(client, task_id, lambda t: t["status"] == "rejected")
    assert client.post(f"/api/tasks/{task_id}/plan/decision", headers=ENG,
                       json={"decision": "approve"}).status_code == 409


def test_auth_and_access_control(api: tuple[TestClient, Runtime]) -> None:
    client, _ = api
    assert client.get("/api/me").status_code == 401
    assert client.get("/api/me", headers={"X-User": "mallory"}).status_code == 401
    assert client.get("/api/workspaces/proc/files", headers=ENG).status_code == 403
    assert client.get("/api/workspaces/nowhere/files", headers=ENG).status_code == 404
    resp = client.post("/api/tasks", headers=ENG, json={"workspace": "plant-a", "text": "Hello",
                                                         "attachments": ["../config/users.yaml"]})
    assert resp.status_code == 403
    offer = client.post("/api/tasks", headers={"X-User": "buyer1"},
                        json={"workspace": "proc", "text": "What is the delivery period in this offer?",
                              "attachments": ["inputs/offer_b.pdf"]}).json()
    assert client.get(f"/api/tasks/{offer['id']}", headers=ENG).status_code == 403
    assert client.get(f"/api/tasks/{offer['id']}", headers={"X-User": "admin1"}).status_code == 403
    files = client.get("/api/workspaces/plant-a/files", headers={"X-User": "dev1"}).json()
    assert files and all(f["label"]["level"] in {"Unclassified", "Restricted"} for f in files)
    upload = client.post("/api/workspaces/plant-a/files", headers={"X-User": "dev1"},
                         files={"file": ("x.md", b"hello", "text/markdown")}, data={"level": "Secret"})
    assert upload.status_code == 403
    client.cookies.set("wb_user", "officer1")
    cookie = client.get("/api/me")
    client.cookies.clear()
    assert cookie.status_code == 200 and cookie.json()["id"] == "officer1"


def test_upload_detects_marking(api: tuple[TestClient, Runtime]) -> None:
    client, rt = api
    body = b"# Pump note\n\nCONFIDENTIAL\n\nThe casing was 5.9 mm."
    resp = client.post("/api/workspaces/plant-a/files", headers=ENG, files={"file": ("pump-note.md", body, "text/markdown")},
                       data={"level": "Restricted"})
    assert resp.status_code == 201, resp.text
    assert resp.json()["label_display"] == "Confidential"
    assert client.post("/api/workspaces/plant-a/files", headers=ENG,
                       files={"file": ("pump-note.md", body, "text/markdown")}).status_code == 409
    event = next(e.event for e in rt.audit.entries() if e.event["type"] == "file.upload")
    assert event["chosen"] == "Restricted" and event["detected"] == "Confidential"


def test_system_endpoints(api: tuple[TestClient, Runtime]) -> None:
    client, _ = api
    health = client.get("/api/health").json()
    assert health["sandboxd"]["status"] == "ok" and health["egressd"]["status"] == "ok"
    before = client.get("/api/egress").json()
    assert before["status"] == "ok" and before["external_connections"] == 0
    test = client.post("/api/egress/test", headers=ENG).json()
    assert test["pass"] and [c["name"] for c in test["checks"]] == ["host_raw_ip", "host_dns", "sandbox"]
    after = client.get("/api/egress").json()
    assert after["blocked_connect_host"] == before["blocked_connect_host"] + 1
    assert after["blocked_connect_sandbox"] == before["blocked_connect_sandbox"] + 1
    assert len(client.get("/api/egress/events").json()["events"]) >= 2
    models = client.get("/api/models", headers=ENG).json()
    assert {m["name"] for m in models["models"]} >= {"qwen3-vl-8b", "granite-3.3-8b"}
    assert models["pool"]["state"]["gpt-oss-20b"] == "asleep"
    assert client.get("/api/audit/tail", headers=ENG).status_code == 403
    assert client.get("/api/audit/tail", headers=OFFICER).status_code == 200
    assert client.get("/api/kb/impact/SOP-MECH-014/5", headers=OWNER).json()["previous"] == "4"
    assert client.get("/api/kb/impact/SOP-MECH-014/9", headers=OWNER).status_code == 404
    assert client.get("/api/templates", headers=ENG).json()["templates"]
    assert client.get("/api/users").status_code == 200
    queued = client.post("/api/tasks", headers=ENG, json={"workspace": "plant-a", "text": "Summarise this vendor contract",
                                                           "attachments": ["inputs/vendor_contract.pdf"]}).json()
    jobs = client.get("/api/jobs", headers=ENG).json()
    job = next(j for j in jobs if j["task_id"] == queued["id"])
    assert client.delete(f"/api/jobs/{job['id']}", headers={"X-User": "buyer1"}).status_code == 403
    assert client.delete(f"/api/jobs/{job['id']}", headers=ENG).status_code == 200
    wait_for(client, queued["id"], lambda t: t["status"] == "cancelled")


def test_ui_pages_have_no_external_urls(api: tuple[TestClient, Runtime]) -> None:
    client, rt = api
    for path in ["/", "/models", "/security", "/t/T12345678", "/t/T12345678/review"]:
        page = client.get(path)
        assert page.status_code == 200
        assert "default-src 'self'" in page.headers["content-security-policy"]
        assert not re.search(r"""(src|href)=["']https?://""", page.text)
    assert client.get("/t/not-a-task").status_code == 404
    for asset in ["/static/app.js", "/static/app.css", "/static/mark.svg", "/favicon.ico"]:
        body = client.get(asset)
        assert body.status_code == 200
        assert not re.search(r"https?://(?!www\.w3\.org)", body.text)
    tpl_dir = rt.settings.root / "workbench" / "ui" / "templates"
    for tpl in tpl_dir.glob("*.html"):
        assert not re.search(r"""(src|href)=["']https?://""", tpl.read_text(encoding="utf-8")), tpl.name


def test_plan_approval_covers_new_drafts(api_default: tuple[TestClient, Runtime]) -> None:
    client, _rt = api_default
    resp = client.post("/api/tasks", headers=ENG, json={"workspace": "plant-a", "text": "Summarise this vendor contract",
                                                         "attachments": ["inputs/vendor_contract.pdf"]})
    task_id = resp.json()["id"]
    wait_for(client, task_id, gate_is("plan"))
    client.post(f"/api/tasks/{task_id}/plan/decision", headers=ENG, json={"decision": "approve"})
    task = wait_for(client, task_id, gate_is("deliverable"))
    action = next(g for g in task["gates"] if g["kind"] == "action")
    assert action["status"] == "approved" and action["decided_by"] == "engineer1"
    assert action["note"] == "approved with the plan"
    assert all(d["relpath"].startswith(f"drafts/{task_id}/") for d in task["deliverables"])
    # A revision overwrites the draft, so it asks again.
    client.post(f"/api/tasks/{task_id}/draft/decision", headers=ENG,
                json={"decision": "reject", "note": "Add the payment terms."})
    task = wait_for(client, task_id, gate_is("action"))
    assert task["pending_gate"]["payload"]["tool"] == "make_docx"


def test_read_only_question_starts_without_a_plan_approval(api_default: tuple[TestClient, Runtime]) -> None:
    client, rt = api_default
    resp = client.post("/api/tasks", headers=ENG, json={"workspace": "plant-a",
                                                         "text": "Which pumps are governed by SOP-MECH-014?"})
    task_id = resp.json()["id"]
    task = wait_for(client, task_id, lambda t: t["status"] == "completed")
    plan_gate = next(g for g in task["gates"] if g["kind"] == "plan")
    assert plan_gate["decided_by"] == "system" and "read-only" in plan_gate["note"]
    assert "P-108B" in task["result"]["answer"]["answer"][0]["text"]
    assert any(e.event["type"] == "gate.plan" and e.event["by"] == "system" for e in rt.audit.entries())


def test_library_lists_files_and_decisions(api: tuple[TestClient, Runtime]) -> None:
    client, _rt = api
    draft = run_to_draft(client, "Draft an approval note for this inspection report", ["inputs/inspection_P108B.pdf"])
    lib = client.get("/api/library", headers=ENG).json()
    entry = next(t for t in lib["tasks"] if t["id"] == draft["id"])
    assert entry["waiting_for"] == "deliverable"
    assert [f["name"] for f in entry["files"]] == ["approval-note.docx"] and not entry["files"][0]["final"]
    kinds = [(d["kind"], d["status"], d["by"]) for d in entry["decisions"]]
    assert ("plan", "approved", "engineer1") in kinds and ("action", "approved", "engineer1") in kinds
    # A user without access to the workspace sees nothing of it.
    assert all(t["workspace"] != "plant-a" for t in client.get("/api/library", headers={"X-User": "buyer1"}).json()["tasks"])


def test_follow_ups_join_the_conversation(api_default: tuple[TestClient, Runtime]) -> None:
    client, _rt = api_default
    first = client.post("/api/tasks", headers=ENG, json={"workspace": "plant-a",
                                                          "text": "What is the design pressure in the data sheet?",
                                                          "attachments": ["inputs/pipe_data.md"]}).json()
    wait_for(client, first["id"], lambda t: t["status"] == "completed")
    second = client.post(f"/api/tasks/{first['id']}/followup", headers=ENG,
                         json={"text": "How many words are in it?"}).json()
    assert second["followup_of"] == first["id"] and second["attachments"] == ["inputs/pipe_data.md"]
    wait_for(client, second["id"], lambda t: t["status"] == "completed")
    # A follow-up of a follow-up still joins the root conversation.
    third = client.post(f"/api/tasks/{second['id']}/followup", headers=ENG, json={"text": "And how many pages?"}).json()
    assert third["followup_of"] == first["id"]
    wait_for(client, third["id"], lambda t: t["status"] == "completed")
    root = client.get(f"/api/tasks/{first['id']}", headers=ENG).json()
    assert [f["id"] for f in root["followups"]] == [second["id"], third["id"]]
    assert "words" in root["followups"][0]["result"]["answer"]["answer"][0]["text"]
    assert root["followups"][0]["label_display"] == "Restricted"


def test_small_talk_and_vague_requests_get_useful_replies(api_default: tuple[TestClient, Runtime]) -> None:
    client, _rt = api_default
    root = client.post("/api/tasks", headers=ENG, json={"workspace": "plant-a", "text": "hello"}).json()
    done = wait_for(client, root["id"], lambda t: t["status"] == "completed")
    reply = done["result"]["reply"]
    assert reply["text"].startswith("Hello!") and "approval notes" in reply["text"]
    assert any(s["label"] == "Summarise the contract" for s in reply["suggestions"])
    assert "answer" not in done["result"] and done["label_display"] == "Restricted"

    turns = {}
    for text in ["summarise any document", "continue", "summarize!"]:
        child = client.post(f"/api/tasks/{root['id']}/followup", headers=ENG, json={"text": text}).json()
        turns[text] = wait_for(client, child["id"], lambda t: t["status"] == "completed")["result"]["reply"]
    ask = turns["summarise any document"]
    assert ask["text"].startswith("Which document should I summarise?")
    assert {s["attachments"][0] for s in ask["suggestions"]} <= {
        "inputs/inspection_P101A_clean.pdf", "inputs/inspection_P101A_injected.pdf", "inputs/inspection_P108B.pdf",
        "inputs/vendor_contract.pdf"}
    assert "summarise any document" in turns["continue"]["text"]
    assert turns["summarize!"]["text"].startswith("Which document should I summarise?")

    # Picking a suggestion runs the real task with that file, in the same conversation.
    pick = next(s for s in ask["suggestions"] if s["attachments"] == ["inputs/vendor_contract.pdf"])
    child = client.post(f"/api/tasks/{root['id']}/followup", headers=ENG,
                        json={"text": pick["text"], "attachments": pick["attachments"]}).json()
    assert child["attachments"] == ["inputs/vendor_contract.pdf"] and child["followup_of"] == root["id"]
    task = wait_for(client, child["id"], gate_is("plan"))
    assert task["plan"]["template"] == "contract_summary"


def test_a_named_file_is_attached_automatically(api_default: tuple[TestClient, Runtime]) -> None:
    client, _rt = api_default
    task = client.post("/api/tasks", headers=ENG, json={"workspace": "plant-a",
                                                         "text": "Summarise the vendor contract"}).json()
    assert task["attachments"] == ["inputs/vendor_contract.pdf"]
    other = client.post("/api/tasks", headers=ENG, json={"workspace": "plant-a",
                                                          "text": "Which pumps are governed by SOP-MECH-014?"}).json()
    assert other["attachments"] == []


def test_a_paragraph_can_be_rewritten_and_must_then_be_accepted(api: tuple[TestClient, Runtime]) -> None:
    client, rt = api
    task = run_to_draft(client, "Draft an approval note for this inspection report", ["inputs/inspection_P108B.pdf"])
    task_id = task["id"]
    draft = client.get(f"/api/tasks/{task_id}/draft", headers=ENG).json()
    d = draft["deliverables"][0]
    body = next(b for b in d["preview"]["blocks"]
                if b["type"] == "p" and b["style"] not in {"Title", "WB Marking", "WB Reference"}
                and any(r["sup"] for r in b["runs"]))
    index = body["index"]
    markers_before = [r["text"] for r in body["runs"] if r["sup"]]

    edited = client.post(f"/api/tasks/{task_id}/draft/edit", headers=ENG,
                         json={"file_id": d["file_id"], "paragraph": index,
                               "text": "The casing was inspected and a reading of 99.9 mm was recorded."}).json()
    page = edited["deliverables"][0]
    written = next(b for b in page["preview"]["blocks"] if b.get("index") == index)
    assert "".join(r["text"] for r in written["runs"] if not r["sup"]).startswith("The casing was inspected")
    # The reference markers the renderer placed are still there.
    assert [r["text"] for r in written["runs"] if r["sup"]] == markers_before
    # A figure typed in by hand has no record behind it, so it is flagged at once.
    assert "99.9 mm" in [f["raw"] for f in page["provenance"]["figures"] if f["status"] == "unsourced"]

    key = f"{d['file_id']}:{index}"
    edit = next(e for e in edited["summary"]["edits"] if e["key"] == key)
    assert edit["by"] == "engineer1" and edit["accepted"] is False
    assert "1 edited paragraph(s) not accepted" in edited["blockers"]
    assert client.post(f"/api/tasks/{task_id}/draft/decision", headers=ENG,
                       json={"decision": "approve"}).status_code == 409

    accepted = client.post(f"/api/tasks/{task_id}/draft/edits/{key}/accept", headers=ENG).json()
    assert accepted["summary"]["edits"][0]["accepted"] is True
    assert "edited paragraph" not in " ".join(accepted["blockers"])
    trail = {e.event["type"] for e in rt.audit.tail(60) if str(e.event.get("type", "")).startswith("draft.edit")}
    assert trail == {"draft.edited", "draft.edit_accepted"}


def test_an_approved_file_cannot_be_rewritten(api: tuple[TestClient, Runtime]) -> None:
    client, _rt = api
    task = run_to_draft(client, "Draft an approval note for this inspection report",
                        ["inputs/inspection_P108B.pdf"])
    task_id = task["id"]
    draft = client.get(f"/api/tasks/{task_id}/draft", headers=ENG).json()
    d = draft["deliverables"][0]
    acknowledged = [c["record_id"] for c in draft["checks"] if c["status"] == "mismatch"]
    assert client.post(f"/api/tasks/{task_id}/draft/decision", headers=ENG,
                       json={"decision": "approve", "acknowledged": acknowledged}).status_code == 200
    wait_for(client, task_id, lambda t: t["status"] == "completed")
    refused = client.post(f"/api/tasks/{task_id}/draft/edit", headers=ENG,
                          json={"file_id": d["file_id"], "paragraph": 4, "text": "Changed after approval."})
    assert refused.status_code == 409
    assert "approved file cannot be edited" in refused.json()["detail"]
