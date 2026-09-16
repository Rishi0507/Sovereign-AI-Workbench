"""Tools, sandbox, documents and knowledge base (PRD phases 3 to 5)."""

from __future__ import annotations

import importlib
import os
import sys
from datetime import date
from pathlib import Path
from typing import Any

import pytest
from docx import Document
from openpyxl import load_workbook
from pptx import Presentation

from workbench.core.errors import JailError, NotFound, PolicyError
from workbench.core.labels import Label, Level
from workbench.documents.dual_read import normalise_field, reconcile
from workbench.documents.pipeline import ReadContext, read_document
from workbench.documents.readers import CompositeReader, FixtureReader, Region, TextPdfReader
from workbench.kb.impact import impact_report
from workbench.llm.heuristic import HeuristicBackend
from workbench.runtime import Runtime
from workbench.tools.render_common import read_embedded_label, restamp
from workbench.tools.sandbox import FakeSandbox, traceback_head
from workbench.workspace import clean_relpath

TOOL_MODULES = ["files", "run_python", "read_document", "search_kb", "check_consistency", "calculate",
                "render_docx", "render_xlsx", "render_pptx", "misc"]


def ctx_for(rt: Runtime, task_id: str = "T1", workspace: str = "plant-a", user: str = "engineer1") -> Any:
    from workbench.core.models import Anchor

    if not rt.ledger.for_task(task_id):
        rt.ledger.add(task_id, "user_input", summary="test", body="test", label=Label(level=Level.RESTRICTED))
        rt.ledger.add(task_id, "attachment", summary="a", body={}, label=Label(level=Level.CONFIDENTIAL),
                      anchor=Anchor(doc="x"))
    return rt.orchestrator.tool_context(
        rt.tasks.get(task_id) if _has_task(rt, task_id) else _fake_state(rt, task_id, workspace, user),
        "s1", "qwen3-vl-8b", "call-1", {"action_approved": True})


def _has_task(rt: Runtime, task_id: str) -> bool:
    try:
        rt.tasks.get(task_id)
        return True
    except NotFound:
        return False


def _fake_state(rt: Runtime, task_id: str, workspace: str, user: str) -> Any:
    from workbench.agent.state import TaskState

    now = rt.clock.now().isoformat()
    state = TaskState(id=task_id, workspace=workspace, user=user, text="t", created_at=now, updated_at=now)
    return rt.tasks.create(state)


# -- workspace jail ---------------------------------------------------------------------------


@pytest.mark.parametrize("bad", ["../x", "/etc/passwd", "inputs/../../x", "inputs/./a", "ｉnputs/a", "a\u202eb",
                                 "C:\\x", "inputs/a*b", ""])
def test_clean_relpath_rejects(bad: str) -> None:
    with pytest.raises(JailError):
        clean_relpath(bad)


def test_jail_rejects_symlinks_and_escapes(rt: Runtime, tmp_path: Path) -> None:
    base = rt.files.ws_root("plant-a")
    outside = tmp_path / "secret.txt"
    outside.write_text("x", encoding="utf-8")
    link = base / "inputs" / "link.txt"
    try:
        os.symlink(outside, link)
    except (OSError, NotImplementedError):
        pytest.skip("symlinks are not available for this user")
    with pytest.raises(JailError):
        rt.files.resolve("plant-a", "inputs/link.txt")


def test_file_store_upload_share_and_final(rt: Runtime) -> None:
    rec = rt.files.upload("plant-a", "notes.md", b"# CONFIDENTIAL notes\nvalue 5 mm", None, "engineer1",
                          detected=Label(level=Level.CONFIDENTIAL))
    assert rec.label.level == Level.CONFIDENTIAL and rec.area == "inputs"
    with pytest.raises(PolicyError):
        rt.files.upload("plant-a", "notes.md", b"again", None, "engineer1")
    with pytest.raises(PolicyError):
        rt.files.upload("plant-a-general", "hi.md", b"x", Label(level=Level.SECRET), "engineer1")
    with pytest.raises(PolicyError):
        rt.files.share(rec.id, "plant-a-general", "engineer1")
    low = rt.files.upload("plant-a", "public.md", b"x", Label(level=Level.RESTRICTED), "engineer1")
    shared = rt.files.share(low.id, "plant-a-general", "engineer1")
    assert shared.workspace == "plant-a-general" and shared.label == low.label
    with pytest.raises(PolicyError):
        rt.files.share(low.id, "proc", "engineer1")
    draft = rt.files.write_draft("plant-a", "d.md", b"draft", Label(level=Level.CONFIDENTIAL), "T1", False)
    with pytest.raises(PolicyError):
        rt.files.write_draft("plant-a", "d.md", b"again", Label(level=Level.CONFIDENTIAL), "T1", False)
    final = rt.files.promote_to_final(draft.id, "engineer1")
    assert final.relpath == "final/d.md" and final.label == draft.label
    with pytest.raises(PolicyError):
        rt.files.promote_to_final(rec.id, "engineer1")


# -- tool registry --------------------------------------------------------------------------


def test_tool_handlers_import_and_run_without_network(rt: Runtime) -> None:
    for name in TOOL_MODULES:
        importlib.import_module(f"workbench.tools.{name}")
    ctx = ctx_for(rt)
    assert rt.tools.execute("list_files", {"area": "inputs"}, ctx).ok
    res = rt.tools.execute("search_kb", {"queries": ["minimum thickness"], "top_k": 3}, ctx)
    assert res.ok and res.records


def test_no_tool_can_change_labels(rt: Runtime) -> None:
    for spec in rt.tools.specs():
        props = spec.input_schema.get("properties", {})
        assert not ({"label", "labels", "classification", "marking", "level"} & set(props)), spec.name
    assert "fetch" not in rt.tools.names() and "http_get" not in rt.tools.names()


def test_tool_validation_and_unknown(rt: Runtime) -> None:
    ctx = ctx_for(rt)
    assert not rt.tools.execute("nope", {}, ctx).ok
    bad = rt.tools.execute("search_kb", {"queries": []}, ctx)
    assert not bad.ok and "invalid arguments" in bad.summary
    assert "(root)" in rt.tools.validate("read_file", {})[0]


def test_write_file_only_under_drafts(rt: Runtime) -> None:
    ctx = ctx_for(rt)
    assert rt.tools.execute("write_file", {"path": "drafts/a.md", "content": "hello"}, ctx).ok
    assert not rt.tools.execute("write_file", {"path": "final/a.md", "content": "x"}, ctx).ok
    ctx.meta["action_approved"] = False
    again = rt.tools.execute("write_file", {"path": "drafts/a.md", "content": "y"}, ctx)
    assert not again.ok and "approval" in again.summary


def test_read_file_records_text(rt: Runtime) -> None:
    ctx = ctx_for(rt)
    res = rt.tools.execute("read_file", {"path": "inputs/pipe_data.md"}, ctx)
    rec = rt.ledger.get(res.records[0])
    assert rec.kind == "ocr_text" and "Design pressure" in rec.body_text()
    assert rec.label.level == Level.RESTRICTED
    assert not rt.tools.execute("read_file", {"path": "../config/users.yaml"}, ctx).ok


def test_calculate_units_and_steps(rt: Runtime) -> None:
    ctx = ctx_for(rt)
    variables = {"P": {"value": 4.5, "unit": "MPa"}, "D": {"value": 219.1, "unit": "mm"},
                 "S": {"value": 138, "unit": "MPa"}, "E": {"value": 1.0, "unit": "dimensionless"},
                 "Y": {"value": 0.4, "unit": "dimensionless"}, "CA": {"value": 1.5, "unit": "mm"}}
    res = rt.tools.execute("calculate", {"expression": "t = P*D/(2*(S*E + P*Y)) + CA", "variables": variables,
                                         "result_unit": "mm"}, ctx)
    assert res.ok
    expected = 4.5 * 219.1 / (2 * (138 * 1.0 + 4.5 * 0.4)) + 1.5
    assert res.body["result"]["value"] == pytest.approx(expected, rel=1e-6)
    assert all("mm" in s["result"] or s["result"] == "" for s in res.body["steps"])
    assert res.body["excel"]["raw_units_consistent"]
    bad = rt.tools.execute("calculate", {"expression": "__import__('os').system('x')", "variables": variables,
                                         "result_unit": "mm"}, ctx)
    assert not bad.ok
    mixed = rt.tools.execute("calculate", {"expression": "L = a + b", "result_unit": "mm",
                                           "variables": {"a": {"value": 1, "unit": "m"},
                                                         "b": {"value": 5, "unit": "mm"}}}, ctx)
    assert mixed.body["result"]["value"] == pytest.approx(1005.0) and not mixed.body["excel"]["raw_units_consistent"]
    wrong = rt.tools.execute("calculate", {"expression": "L = a + b", "result_unit": "mm",
                                           "variables": {"a": {"value": 1, "unit": "m"},
                                                         "b": {"value": 5, "unit": "bar"}}}, ctx)
    assert not wrong.ok and "unit" in wrong.summary


def test_renderers_mark_and_keep_formulas(rt: Runtime) -> None:
    ctx = ctx_for(rt)
    note = {"title": "Note", "subject": "S", "summary": [{"text": "Reading 5.6 mm [R-T1-1]."}],
            "findings": [], "consistency_findings": [{"rule": "r", "status": "mismatch", "text": "x"}],
            "recommendation": [{"text": "Stop."}]}
    doc = rt.tools.execute("make_docx", {"template": "approval_note", "data": note}, ctx)
    path = rt.files.open_path(doc.files[0])
    d = Document(str(path))
    assert d.sections[0].header.paragraphs[0].text == "CONFIDENTIAL"
    assert d.sections[0].footer.paragraphs[0].text == "CONFIDENTIAL"
    assert "classification=CONFIDENTIAL" in d.core_properties.keywords
    assert any(r.font.superscript for p in d.paragraphs for r in p.runs)
    assert read_embedded_label(path) == Label(level=Level.CONFIDENTIAL)
    calc = rt.tools.execute("calculate", {"expression": "L = a + b", "result_unit": "mm",
                                          "variables": {"a": {"value": 1, "unit": "m", "record": "R-T1-1"},
                                                        "b": {"value": 5, "unit": "mm"}}}, ctx)
    xl = rt.tools.execute("make_xlsx", {"kind": "calc_sheet", "data": {"calc": calc.body}}, ctx)
    wb = load_workbook(str(rt.files.open_path(xl.files[0])))
    ws = wb["Calculation"]
    formulas = [c.value for row in ws.iter_rows() for c in row if isinstance(c.value, str) and c.value.startswith("=")]
    assert any("ROUND" in f for f in formulas) and ws.oddHeader.center.text == "CONFIDENTIAL"
    assert ws["C6"].comment is not None and "R-T1-1" in ws["C6"].comment.text
    deck = rt.tools.execute("make_pptx", {"data": {"title": "Board", "slides": [{"title": "One",
                                                                              "bullets": ["a", "b"]}]}}, ctx)
    prs = Presentation(str(rt.files.open_path(deck.files[0])))
    assert len(prs.slides) == 2
    for slide in prs.slides:
        assert any(s.name == "WB Marking" and s.text_frame.text == "CONFIDENTIAL" for s in slide.shapes)
    restamp(path, Label(level=Level.RESTRICTED))
    assert Document(str(path)).sections[0].header.paragraphs[0].text == "RESTRICTED"
    assert read_embedded_label(path) == Label(level=Level.RESTRICTED)


def test_recall_is_limited_to_task(rt: Runtime) -> None:
    ctx = ctx_for(rt, "T1")
    other = rt.ledger.add("T2", "ocr_text", summary="x", body="secret", label=Label.lowest())
    assert not rt.tools.execute("recall", {"id": other.id}, ctx).ok
    own = rt.ledger.for_task("T1")[0]
    assert rt.tools.execute("recall", {"id": own.id}, ctx).body["text"] == "test"


# -- sandbox --------------------------------------------------------------------------------


def test_fake_sandbox_failure_success_and_network(tmp_path: Path) -> None:
    reports: list[dict[str, Any]] = []

    class Reporter:
        def report(self, event: dict[str, Any]) -> None:
            reports.append(event)

    sb = FakeSandbox(reporter=Reporter())
    (tmp_path / "bad.py").write_text("d = {}\nd['pressure']\n", encoding="utf-8")
    res = sb.run(tmp_path, "bad.py", 30, run_id="r1")
    assert res.exit_code == 1 and res.traceback_head[0].startswith("Traceback")
    assert "KeyError" in res.stderr_tail
    (tmp_path / "ok.py").write_text("open('out.txt','w').write('x')\nprint('{\"n\": 1}')\n", encoding="utf-8")
    ok = sb.run(tmp_path, "ok.py", 30, run_id="r2")
    assert ok.exit_code == 0 and [f.path for f in ok.new_files] == ["out.txt"]
    (tmp_path / "net.py").write_text(
        "import socket\ntry:\n    socket.create_connection(('1.1.1.1', 443), 5)\nexcept OSError as e:\n"
        "    print('blocked', e.errno)\n", encoding="utf-8")
    net = sb.run(tmp_path, "net.py", 30, run_id="r3")
    assert net.exit_code == 0 and len(net.net_attempts) == 1 and net.net_attempts[0].addr == "1.1.1.1:443"
    assert reports == [{"origin": "sandbox", "addr": "1.1.1.1:443", "run_id": "r3", "source": "probe"}]
    (tmp_path / "dns.py").write_text("import socket\ntry:\n    socket.getaddrinfo('example.com', 443)\n"
                                     "except OSError:\n    pass\n", encoding="utf-8")
    assert len(sb.run(tmp_path, "dns.py", 30, run_id="r4").net_attempts) == 1
    (tmp_path / "slow.py").write_text("import time\ntime.sleep(10)\n", encoding="utf-8")
    slow = sb.run(tmp_path, "slow.py", 1, run_id="r5")
    assert slow.timed_out
    with pytest.raises(Exception):
        sb.run(tmp_path, "../bad.py", 5)
    env_leak = tmp_path / "env.py"
    env_leak.write_text("import os, sys\nprint(sorted(k for k in os.environ if k.startswith('WB_')))\n"
                        "print(sys.flags.ignore_environment)\n", encoding="utf-8")
    leak = sb.run(tmp_path, "env.py", 10, run_id="r6")
    assert leak.stdout_tail.splitlines() == ["[]", "1"]


def test_traceback_head_uses_last_traceback() -> None:
    text = "Traceback (most recent call last):\n a\nX\nTraceback (most recent call last):\n b\nY: z"
    assert traceback_head(text, 2) == ["Traceback (most recent call last):", " b"]


# -- documents --------------------------------------------------------------------------------


def _read_ctx(rt: Runtime, task: str = "TD") -> ReadContext:
    return ReadContext(task_id=task, ledger=rt.ledger, backend=HeuristicBackend(rt.ledger), vlm_model="qwen3-vl-8b",
                       policy=rt.policy.policy, evidence_dir=rt.evidence_dir(task), produced_by="test")


def test_dual_read_medium_and_uncertain(rt: Runtime) -> None:
    path = rt.files.resolve("plant-a", "inputs/inspection_P108B.pdf")
    result = read_document(path, Label(level=Level.CONFIDENTIAL), _read_ctx(rt))
    tag = next(r for r in result.records if r.kind == "vlm_read" and r.body["field_kind"] == "tag")
    assert tag.body["value"] == "P-108B" and tag.confidence == "medium" and tag.body["status"] == "resolved"
    assert tag.body["ocr_value"] == "P-1O8B" and tag.anchor and tag.anchor.page == 1 and tag.anchor.region
    assert tag.body["crop"] and (rt.evidence_dir("TD") / tag.body["crop"]).is_file()
    wo = next(r for r in result.records if r.kind == "vlm_read" and r.body["region"] == "p3-wo")
    assert wo.confidence == "uncertain" and len(result.uncertain) == 1
    agreed = next(r for r in result.records if r.kind == "vlm_read" and r.body["region"] == "p1-date")
    assert agreed.confidence == "high"
    assert all(r.label.level == Level.CONFIDENTIAL for r in result.records)
    assert all(r.anchor is not None and r.anchor.doc == "inspection_P108B.pdf" for r in result.records)
    ocr = [r for r in result.records if r.kind == "ocr_text"]
    assert len(ocr) == 3 and ocr[1].body_text().count("|") > 10


def test_dual_read_neither_is_uncertain() -> None:
    region = Region(id="r", bbox=(0, 0, 1, 1), field_kind="tag", ocr_value="P-1O8B", ocr_conf=0.5, needs_vlm=True,
                    vlm_value="P-109B", vlm_value_zoomed="P-107B")
    out = reconcile(region, HeuristicBackend(), "m")
    assert out.status == "uncertain" and out.value.confidence == "uncertain"
    assert normalise_field("stamp", "[ STAMP ]") == "present"
    assert normalise_field("quantity", "5.6 mm") == "5.6 millimeter"
    assert normalise_field("date", "not a date") is None


def test_hindi_fixture_and_marking(rt: Runtime) -> None:
    path = rt.files.resolve("plant-a", "inputs/inspection_hindi_mixed.ocr.json")
    result = read_document(path, Label(level=Level.RESTRICTED), _read_ctx(rt, "TH"))
    assert result.pages[0].script == "mixed"
    assert result.label.level == Level.CONFIDENTIAL, "marking detected in the page text raises the label"
    tag = next(r for r in result.records if r.kind == "vlm_read" and r.body["field_kind"] == "tag")
    assert tag.body["value"] == "E-201" and tag.confidence == "medium"


def test_text_pdf_reader(rt: Runtime) -> None:
    pages = TextPdfReader().read(rt.files.resolve("proc", "inputs/offer_a.pdf"))
    assert len(pages) == 3 and not pages[0].scanned
    kv = next(t for t in pages[0].tables if t.id.endswith("kv"))
    assert ["Total price", "INR 48,50,000"] in kv.rows
    assert any(r.field_kind == "quantity" and "48,50,000" in r.ocr_value for r in pages[0].regions)
    assert FixtureReader().can_read(rt.files.resolve("plant-a", "inputs/inspection_P108B.pdf"))
    assert CompositeReader().engine_for(rt.files.resolve("proc", "inputs/offer_a.pdf")) == "pymupdf"


# -- knowledge base ---------------------------------------------------------------------------


def test_revision_aware_retrieval_and_clause_diff(rt: Runtime) -> None:
    user, ws = rt.policy.user("engineer1"), rt.policy.workspace("plant-a")
    q = ["minimum casing wall thickness centrifugal pump"]
    old = rt.kb.search(q, 3, date(2023, 6, 1), user, ws, rt.policy, today=date(2026, 9, 15))
    new = rt.kb.search(q, 3, date(2026, 9, 15), user, ws, rt.policy, today=date(2026, 9, 15))
    top_old = next(h for h in old.hits if "4.3" in h.chunk.clause_ids)
    top_new = next(h for h in new.hits if "4.3" in h.chunk.clause_ids)
    assert top_old.chunk.revision == "4" and "5.5 mm" in top_old.chunk.text and top_old.currency_warning
    assert top_new.chunk.revision == "5" and "6.0 mm" in top_new.chunk.text and top_new.currency_warning is None
    diff = {d.clause: d for d in rt.kb.clause_diff("SOP-MECH-014", "5")}
    assert diff["4.3"].status == "amended"
    assert (diff["4.3"].limits[0].old, diff["4.3"].limits[0].new, diff["4.3"].limits[0].unit) == (5.5, 6.0, "millimeter")
    assert diff["4.4"].status == "withdrawn" and diff["4.5"].status == "added" and diff["4.1"].status == "unchanged"


def test_impact_report(rt: Runtime) -> None:
    report = impact_report(rt.kb, "SOP-MECH-014", "5")
    assert {n["note"] for n in report["notes"]} == {"AN-2023-017", "AN-2024-004"}
    assert "P-108B" in report["equipment"] and "E-201" not in report["equipment"]
    assert "Nothing has been changed automatically" in report["markdown"]


def test_acl_and_label_filters(rt: Runtime) -> None:
    today = date(2026, 9, 15)
    eng, plant = rt.policy.user("engineer1"), rt.policy.workspace("plant-a")
    res = rt.kb.search(["evaluation of commercial offers weights"], 10, today, eng, plant, rt.policy)
    assert all(h.chunk.doc_number != "PROC-POL-007" for h in res.hits)
    assert res.filtered_out.get("acl", 0) >= 1
    buyer, proc = rt.policy.user("buyer1"), rt.policy.workspace("proc")
    res2 = rt.kb.search(["evaluation of commercial offers weights"], 5, today, buyer, proc, rt.policy)
    assert any(h.chunk.doc_number == "PROC-POL-007" for h in res2.hits)
    dev, _ = rt.policy.user("dev1"), None
    res3 = rt.kb.search(["past approval notes P-108B findings"], 10, today, dev, plant, rt.policy)
    assert all(h.chunk.label.level <= Level.RESTRICTED for h in res3.hits)
    assert res3.filtered_out.get("label", 0) >= 1


def test_graph_expansion_reaches_linked_documents(rt: Runtime) -> None:
    eng, plant = rt.policy.user("engineer1"), rt.policy.workspace("plant-a")
    res = rt.kb.search(["standby unit for P-108B"], 8, date(2026, 9, 15), eng, plant, rt.policy)
    docs = {h.chunk.doc_number for h in res.hits}
    assert "PID-CW-003" in docs
    assert any("graph" in h.via for h in res.hits)
    assert [n.props["value"] for n in res.graph_facts] == [6.8, 6.3]
    past = rt.kb.search(["standby unit for P-108B"], 8, date(2023, 1, 1), eng, plant, rt.policy)
    assert [n.props["value"] for n in past.graph_facts] == [6.8]
    neighbours = {n.kind for n, _ in rt.kb.graph.neighbours("P108B")}
    assert {"class", "vendor", "po", "document", "inspection"} <= neighbours
    assert rt.kb.graph.tags_of_class("centrifugal_pump") == ["P-101A", "P-101B", "P-108A", "P-108B"]


def test_workspace_document_index_is_scoped(rt: Runtime) -> None:
    ctx = ctx_for(rt, "TW")
    assert rt.tools.execute("read_document", {"path": "inputs/vendor_contract.pdf"}, ctx).ok
    assert rt.kb.has_document("plant-a", "vendor_contract.pdf")
    eng = rt.policy.user("engineer1")
    hits = rt.kb.search(["liquidated damages"], 3, date(2026, 9, 15), eng, rt.policy.workspace("plant-a"),
                        rt.policy, doc_filter="plant-a/vendor_contract.pdf").hits
    assert hits and all(h.chunk.workspace == "plant-a" for h in hits)
    officer = rt.policy.user("officer1")
    other = rt.kb.search(["liquidated damages"], 10, date(2026, 9, 15), officer, rt.policy.workspace("proc"),
                         rt.policy).hits
    assert all(h.chunk.workspace != "plant-a" for h in other)


def test_probe_is_standard_library_only() -> None:
    probe = Path(sys.modules["workbench"].__file__).parent / "security" / "probe"
    for f in probe.glob("*.py"):
        text = f.read_text(encoding="utf-8")
        assert "import workbench" not in text and "from workbench" not in text
