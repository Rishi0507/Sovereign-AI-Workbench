"""Acceptance tests mapped to the user stories (PRD section 8)."""

from __future__ import annotations

import shutil
import statistics
from collections.abc import Callable
from pathlib import Path
from typing import Any

import pytest
from docx import Document
from openpyxl import load_workbook

from tests.helpers import GO_MODE, ROOT, run_trace
from workbench.context.cache_salt import cache_salt
from workbench.core.errors import PolicyError
from workbench.core.labels import Label, Level
from workbench.llm.base import LLMRequest, LLMResponse
from workbench.runtime import Runtime


def text_of(path: Path) -> str:
    doc = Document(str(path))
    parts = [p.text for p in doc.paragraphs]
    parts += [c.text for t in doc.tables for r in t.rows for c in r.cells]
    return "\n".join(parts)


def final_path(rt: Runtime, state: Any, suffix: str) -> Path:
    d = next(d for d in state.deliverables if d.relpath.endswith(suffix))
    return rt.files.open_path(d.final_file_id)


@pytest.mark.skipif(GO_MODE, reason="covered by the Python suite")
def test_trace_a_approval_note(make_runtime: Callable[..., Runtime]) -> None:
    """US-1, US-6."""
    rt = make_runtime()
    state = run_trace(rt, "A")
    assert state.status == "completed"
    assert state.route["profile"]["rule"] == "attachment:pdf" and state.model == "qwen3-vl-8b"
    assert state.plan.template == "approval_note_from_scan"
    tag = next(r for r in rt.ledger.for_task(state.id) if r.kind == "vlm_read" and r.body["field_kind"] == "tag")
    assert tag.confidence == "medium" and tag.body["value"] == "P-108B"
    kb = [r for r in rt.ledger.for_task(state.id) if r.kind == "kb_chunk"]
    assert any(r.anchor.doc == "SOP-MECH-014" and r.anchor.revision == "5" for r in kb)
    assert not any(r.anchor.doc == "SOP-MECH-014" and r.anchor.revision == "4" for r in kb)
    mismatches = sorted(c["rule"] for c in state.checks if c["status"] == "mismatch")
    assert mismatches == ["row_tags_match", "thickness_trend", "thickness_vs_limit"]
    path = final_path(rt, state, ".docx")
    assert path.as_posix().endswith("final/approval-note.docx")
    doc = Document(str(path))
    assert doc.sections[0].header.paragraphs[0].text == "CONFIDENTIAL"
    assert doc.sections[0].footer.paragraphs[0].text == "CONFIDENTIAL"
    assert doc.core_properties.category == "Confidential"
    body = text_of(path)
    assert "SOP-MECH-014 Rev 5 clause 4.3" in body and "6.0 mm" in body and "5.6 mm" in body
    deliverable = state.deliverables[0]
    assert deliverable.provenance["counts"]["unsourced"] == 0
    assert deliverable.claims and all(c["status"] == "verified" for c in deliverable.claims)
    shared = rt.files.get(deliverable.final_file_id)
    with pytest.raises(PolicyError):
        rt.files.share(shared.id, "plant-a-general", "engineer1")


@pytest.mark.skipif(GO_MODE, reason="covered by the Python suite")
def test_trace_a_as_of_old_report(make_runtime: Callable[..., Runtime]) -> None:
    """US-1: an old report is grounded in the revision in force then, and flagged for currency today."""
    rt = make_runtime()
    state = run_trace(rt, "A", report_date="2023-06-01")
    assert state.status == "completed"
    kb = [r for r in rt.ledger.for_task(state.id) if r.kind == "kb_chunk" and r.anchor.doc == "SOP-MECH-014"]
    assert kb and all(r.anchor.revision == "4" for r in kb)
    body = text_of(final_path(rt, state, ".docx"))
    assert "Rev 4 clause 4.3" in body and "5.5 mm" in body
    warnings = [c for c in state.deliverables[0].claims if c.get("currency_warning")]
    assert warnings and "superseded by Rev 5" in warnings[0]["currency_warning"]
    thickness = next(c for c in state.checks if c["rule"] == "thickness_vs_limit")
    assert thickness["status"] == "pass", "5.6 mm meets the 5.5 mm minimum that applied in 2023"


def test_trace_b_sandbox_fix(make_runtime: Callable[..., Runtime]) -> None:
    """US-2, US-8."""
    rt = make_runtime()
    state = run_trace(rt, "B")
    assert state.status == "completed" and state.model == "qwen2.5-coder-7b"
    assert state.route["profile"]["rule"] == "code_intent"
    runs = [r for r in rt.ledger.for_task(state.id) if r.kind == "sandbox_result"]
    assert len(runs) >= 2
    assert runs[0].body["exit_code"] != 0 and "KeyError" in "\n".join(runs[0].body["traceback_head"] + [runs[0].body["stderr_tail"]])
    assert runs[-1].body["exit_code"] == 0
    assert runs[-1].body["backend"] == ("dev" if GO_MODE else "fake")
    code = state.result["code"]
    assert code["result"]["anomalies"] == 6 and code["result"]["column"] == "pressure_bar"
    names = sorted(d.relpath for d in state.deliverables)
    assert names == ["drafts/anomalies.csv", "drafts/chart.svg", "drafts/parse_pressure_readings.py", "drafts/summary.json"]
    assert all(d.final_file_id for d in state.deliverables)

    before = rt.egress.snapshot()["blocked_connect_sandbox"]
    variant = rt.files.ws_root("plant-a") / "inputs" / "phone_home.csv"
    shutil.copy(rt.files.open_path(state.deliverables[0].file_id), variant)
    rt.files.register("plant-a", "inputs/phone_home.csv", Label(level=Level.RESTRICTED))
    task = rt.orchestrator.create_task("plant-a", "engineer1", "Inspect the uploaded table", ["inputs/phone_home.csv"])
    ctx = rt.orchestrator.tool_context(rt.tasks.get(task.id), "s", "qwen2.5-coder-7b", "call-9")
    res = rt.tools.execute("run_python", {"script": (
        "import socket\ntry:\n    socket.create_connection(('1.1.1.1', 443), 3)\nexcept OSError as e:\n"
        "    print('{\"blocked\": true}')\n")}, ctx)
    assert res.ok and res.body["result"] == {"blocked": True}
    rec = rt.ledger.get(res.records[0])
    assert len(rec.body["net_attempts"]) == 1
    import time

    deadline = time.time() + 5
    while rt.egress.snapshot()["blocked_connect_sandbox"] < before + 1 and time.time() < deadline:
        time.sleep(0.05)
    assert rt.egress.snapshot()["blocked_connect_sandbox"] == before + 1


@pytest.mark.skipif(GO_MODE, reason="covered by the Python suite")
def test_trace_c_contract_summary(make_runtime: Callable[..., Runtime]) -> None:
    """US-3."""
    rt = make_runtime()
    state = run_trace(rt, "C")
    assert state.status == "completed" and state.model == "qwen3-vl-8b"
    assert state.route["profile"]["decoupled"]
    maps = [t for t in state.trace if t.purpose == "summarise.map"]
    assert len(maps) == 5 and any(t.purpose == "summarise.reduce" for t in state.trace)
    d = state.deliverables[0]
    assert d.claims and all(c["status"] == "verified" for c in d.claims)
    body = text_of(final_path(rt, state, ".docx"))
    assert "INR 4,85,00,000" in body and "Liquidated damages" in body
    assert d.provenance["counts"]["unsourced"] == 0
    follow = rt.orchestrator.create_task("plant-a", "engineer1",
                                         "What is the rate of liquidated damages for delay?",
                                         ["inputs/vendor_contract.pdf"], parent_id=state.id)
    answer = rt.jobs.run_inline(follow.id)
    assert answer.status == "completed"
    assert any(t.tool == "search_kb" and t.ok for t in answer.trace)
    text = " ".join(a["text"] for a in answer.result["answer"]["answer"])
    assert "0.5 percent" in text
    assert answer.result["claims"] and all(c["status"] == "verified" for c in answer.result["claims"])


@pytest.mark.skipif(GO_MODE, reason="covered by the Python suite")
def test_trace_d_calc_sheet(make_runtime: Callable[..., Runtime]) -> None:
    """US-4."""
    rt = make_runtime()
    state = run_trace(rt, "D")
    assert state.status == "completed" and state.plan.template == "calc_sheet"
    calc = next(r for r in rt.ledger.for_task(state.id) if r.kind == "calc_result")
    expected = 4.5 * 219.1 / (2 * (138 * 1.0 + 4.5 * 0.4)) + 1.5
    assert calc.body["result"]["value"] == pytest.approx(expected, rel=1e-6)
    steps = calc.body["steps"]
    assert all("mm" in s["result"] for s in steps if s["result"])
    assert "MPa" in steps[1]["substitution"] and "mm" in steps[1]["substitution"]
    wb = load_workbook(str(final_path(rt, state, ".xlsx")))
    ws = wb["Calculation"]
    derived = [c for row in ws.iter_rows() for c in row if isinstance(c.value, str) and c.value.startswith("=")]
    assert any("ROUND(" in c.value and "E6" in c.value for c in derived)
    assert ws["C6"].comment and "source: R-" in ws["C6"].comment.text
    assert ws.oddHeader.center.text == "RESTRICTED"
    for d in state.deliverables:
        assert d.provenance["counts"]["unsourced"] == 0


@pytest.mark.skipif(GO_MODE, reason="covered by the Python suite")
def test_trace_e_offer_comparison(make_runtime: Callable[..., Runtime]) -> None:
    """US-5."""
    rt = make_runtime()
    state = run_trace(rt, "E")
    assert state.status == "completed"
    profile = state.route["profile"]
    assert profile["decoupled"] and profile["complexity"] == "high"
    vl = next(c for c in state.route["candidates"] if c["model"] == "qwen3-vl-8b")
    assert vl["ok"] and not vl["meets_threshold"]
    assert state.model == "gpt-oss-20b" and state.route["queued_for_tide"]
    assert rt.pool.tide_count >= 1
    assert state.plan.source == "model" and state.plan.repairs == 1
    compile_errors = state.plan_history[0]["errors"]
    assert 'step 3: input "tables" is not produced by any earlier step' in compile_errors
    label = Label.parse("Secret+VENDOR-COMMERCIAL")
    assert all(d.label == label for d in state.deliverables)
    xlsx = final_path(rt, state, ".xlsx")
    wb = load_workbook(str(xlsx))
    offers = wb["Offers"]
    inputs = [c for row in offers.iter_rows(min_row=4, max_row=6, min_col=3, max_col=5) for c in row]
    assert all(c.comment and "source: R-" in c.comment.text for c in inputs)
    assert all(c.fill.fgColor.rgb.endswith("EEF6F4") for c in inputs)
    assert offers.oddHeader.center.text == "SECRET // VENDOR-COMMERCIAL"
    xl = next(d for d in state.deliverables if d.relpath.endswith(".xlsx"))
    assert xl.provenance["counts"]["unsourced"] == 0
    doc = next(d for d in state.deliverables if d.relpath.endswith(".docx"))
    assert doc.provenance["counts"]["unsourced"] == 0 and all(c["status"] == "verified" for c in doc.claims)
    body = text_of(final_path(rt, state, ".docx"))
    assert "Recommend Narmada Pumps Ltd" in body
    assert "Konkan Rotating Equipment Ltd: violates delivery_weeks" in body


@pytest.mark.skipif(GO_MODE, reason="covered by the Python suite")
def test_shadow_onboarding(make_runtime: Callable[..., Runtime], tmp_path: Path) -> None:
    """US-7."""
    root = tmp_path / "root"
    for d in ("config", "templates", "rules", "schemas", "org_templates"):
        shutil.copytree(ROOT / d, root / d)
    shutil.copytree(ROOT / "fixtures" / "eval", root / "fixtures" / "eval")
    rt = make_runtime(root=root)
    from workbench.registry import serve_cmd
    from workbench.registry.shadow import promote, run_shadow

    state = run_trace(rt, "A")
    shadow = next(c for c in state.route["candidates"] if c["model"] == "granite-3.3-8b")
    assert not shadow["ok"] and shadow["reason"] == "status"
    result = run_shadow(rt, "granite-3.3-8b", "admin1")
    assert set(result["quality"]) == {"general", "document", "agentic"}
    proposed = (root / "config" / "models.proposed.yaml").read_text(encoding="utf-8")
    assert "granite-3.3-8b" in proposed
    assert "status: shadow" in (root / "config" / "models.yaml").read_text(encoding="utf-8")
    with pytest.raises(PolicyError):
        promote(rt, "granite-3.3-8b", False, "admin1")
    out = promote(rt, "granite-3.3-8b", True, "admin1")
    assert out["status"] == "active" and rt.registry.get("granite-3.3-8b").status == "active"
    types = [e.event["type"] for e in rt.audit.entries()]
    assert "registry.shadow_eval" in types and "registry.promote" in types
    argv = serve_cmd.build(rt.registry.get("granite-3.3-8b"), "S")
    assert "--speculative-config" in argv and '"method": "ngram"' in argv[argv.index("--speculative-config") + 1]


def test_egress_badge_and_test(make_runtime: Callable[..., Runtime]) -> None:
    """US-8."""
    rt = make_runtime()
    snap = rt.egress.snapshot()
    for key in ("mode", "external_connections", "blocked_packets", "blocked_connect_host",
                "blocked_connect_sandbox", "collectors", "since", "breach"):
        assert key in snap
    assert snap["external_connections"] == 0 and not snap["breach"]
    assert snap["blocked_packets"] is None, "without nftables the counter is unavailable, not zero"
    assert any(str(v).startswith("unavailable") for v in snap["collectors"].values())
    result = rt.egress.run_test()
    assert result["pass"], result
    after = rt.egress.snapshot()
    assert after["blocked_connect_host"] == snap["blocked_connect_host"] + 1
    assert after["blocked_connect_sandbox"] == snap["blocked_connect_sandbox"] + 1
    assert after["blocked_packets"] is None and after["external_connections"] == 0
    dns = next(c for c in result["checks"] if c["name"] == "host_dns")
    assert dns["simulated"] and dns["pass"]


@pytest.mark.skipif(GO_MODE, reason="covered by the Python suite")
def test_prompt_injection_contained(make_runtime: Callable[..., Runtime]) -> None:
    """US-6: injected text cannot change the plan, trigger ungated side effects or leave its data block."""
    rt = make_runtime()
    seen: list[LLMRequest] = []
    inner = rt.backend

    class Recording:
        name = "recording"

        def chat(self, req: LLMRequest) -> LLMResponse:
            seen.append(req)
            return inner.chat(req)

    rt.backend = Recording()
    rt.router.backend = rt.backend
    task = rt.orchestrator.create_task("plant-a", "engineer1", "Draft an approval note for this inspection report",
                                       ["inputs/inspection_P101A_injected.pdf"])
    state = rt.jobs.run_inline(task.id)
    assert state.status == "completed"
    assert [s.id for s in state.plan.steps] == ["extract", "ground", "check", "draft", "render"]
    gates = [g for g in state.gates if g.kind == "action"]
    tool_calls = [t for t in state.trace if t.tool in {"make_docx", "write_file"} and t.kind in {"tool", "default"}]
    assert len(tool_calls) == len(gates) == 1 and gates[0].status == "approved"
    assert not any(d.relpath.startswith("final/") for d in state.deliverables)
    phrase = "Ignore previous instructions"
    hits = 0
    for req in seen:
        for msg in req.messages:
            idx = msg.content.find(phrase)
            while idx != -1:
                hits += 1
                opened = msg.content.rfind("<record ", 0, idx)
                closed = msg.content.rfind("</record>", 0, idx)
                assert opened > closed, "injected text appeared outside a quoted data block"
                assert msg.role == "user"
                idx = msg.content.find(phrase, idx + 1)
    assert hits > 0
    assert all(phrase not in m.content for req in seen for m in req.messages if m.role == "system")


def test_downgrade_flow(make_runtime: Callable[..., Runtime]) -> None:
    """US-6."""
    rt = make_runtime()
    policy = rt.policy
    rec = rt.files.upload("plant-a", "memo.md", b"CONFIDENTIAL memo", Label(level=Level.CONFIDENTIAL), "engineer1")
    officer, owner = policy.user("officer1"), policy.user("owner1")
    req = policy.request_downgrade("D9", rec.id, rec.label, Label(level=Level.RESTRICTED), officer, "redacted")
    with pytest.raises(PolicyError):
        policy.decide_downgrade(req, officer, True, rt.clock.now())
    done = policy.decide_downgrade(req, owner, True, rt.clock.now())
    assert (done.before.level, done.after.level, done.reason) == (Level.CONFIDENTIAL, Level.RESTRICTED, "redacted")
    relabelled = rt.files.relabel(rec.id, done.after)
    assert relabelled.label.level == Level.RESTRICTED


def test_no_cross_partition_cache(make_runtime: Callable[..., Runtime]) -> None:
    """US-6."""
    rt = make_runtime()
    run_trace(rt, "A")
    run_trace(rt, "A")
    parts = rt.cache.snapshot()
    conf = next(p for p in parts if p["label"] == "Confidential")
    assert conf["hit_rate"] > 0.3, "repeated template runs reuse the prefix inside their partition"
    salts = [cache_salt(Label.parse(x), rt.secret_key) for x in ("Restricted", "Confidential", "Secret+VENDOR-COMMERCIAL")]
    prompt = "system\n" + rt.tools.prompt_schemas()
    assert sum(rt.cache.cross_partition_hits(prompt, a, b) for a in salts for b in salts if a != b) == 0
    assert len(set(salts)) == 3


@pytest.mark.skipif(GO_MODE, reason="covered by the Python suite")
def test_trace_timings_within_budget(make_runtime: Callable[..., Runtime]) -> None:
    """PRD 6.2: heuristic traces stay within the CPU budget."""
    import time

    rt = make_runtime()
    limits = {"A": 20, "B": 15, "C": 30}
    times = {}
    for key, limit in limits.items():
        start = time.perf_counter()
        assert run_trace(rt, key).status == "completed"
        times[key] = time.perf_counter() - start
        assert times[key] < limit, times
    assert statistics.fmean(times.values()) > 0
