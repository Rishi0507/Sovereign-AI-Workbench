"""Number provenance: every figure in a deliverable must resolve to evidence (README section 4.6.2)."""

from __future__ import annotations

import json
import math
import re
from collections.abc import Iterable
from pathlib import Path
from typing import Any, Literal

from pydantic import BaseModel, Field

from workbench.core.ledger import Ledger
from workbench.core.models import COMPUTE_KINDS, LedgerRecord
from workbench.core.normalise import (
    DATE_RE,
    QUANTITY_RE,
    TAG_RE,
    canonical_unit,
    parse_number,
    to_base,
    units_compatible,
)

PROVENANCE_KINDS = frozenset({"ocr_text", "vlm_read", "kb_chunk", "graph_fact", "sandbox_result", "calc_result"})
MARKER_RE = re.compile(r"\[(R-[A-Za-z0-9]+-\d+(?:\s*,\s*R-[A-Za-z0-9]+-\d+)*)\]")
REF_RE = re.compile(r"\[(\d+(?:\s*,\s*\d+)*)\]")
MASKS = [
    re.compile(r"R-[A-Za-z0-9]+-\d+"),
    re.compile(r"\bS/N\s*\w+", re.IGNORECASE),
    re.compile(r"\b[A-Z]{1,6}(?:[-/][A-Z0-9]+){1,6}\b"),
    re.compile(r"\bRev(?:ision)?\.?\s*[A-Z0-9]+\b", re.IGNORECASE),
    re.compile(r"\b(?:clauses?|sections?|§)\s*\d+(?:\.\d+)*(?:\s*(?:,|and)\s*\d+(?:\.\d+)*)*", re.IGNORECASE),
    re.compile(r"\b(?:p|pp|page|pages)\.?\s*\d+(?:\s*[-,]\s*\d+)*\b", re.IGNORECASE),
    re.compile(r"\b(?:row|table|step|figure|slide)\s+\d+\b", re.IGNORECASE),
    re.compile(r"^\s*\d+(?:\.\d+)*\.?\s+(?=[A-Z])"),
    re.compile(r"\bT\d+\b"),
]

Status = Literal["sourced", "derived", "unsourced"]


class Figure(BaseModel):
    id: str
    raw: str
    value: float
    unit: str | None = None
    location: str
    context: str = ""
    cited: list[str] = Field(default_factory=list)
    status: Status = "unsourced"
    record_id: str | None = None
    chain: list[str] = Field(default_factory=list)
    resolution: Literal["auto", "linked", "corrected", "confirmed"] = "auto"
    formula: str | None = None
    note: str | None = None


class ProvenanceReport(BaseModel):
    file: str
    figures: list[Figure]

    @property
    def counts(self) -> dict[str, int]:
        out = {"sourced": 0, "derived": 0, "unsourced": 0}
        for f in self.figures:
            out[f.status] += 1
        return out

    def unresolved(self) -> list[Figure]:
        return [f for f in self.figures if f.status == "unsourced" and f.resolution == "auto"]


def mask(text: str) -> str:
    def blank(m: re.Match[str]) -> str:
        return " " * len(m.group(0))

    for pattern in (TAG_RE, DATE_RE, *MASKS):
        text = pattern.sub(blank, text)
    return text


def numbers_in(text: str) -> list[tuple[str, float, str | None]]:
    masked = mask(text)
    out: list[tuple[str, float, str | None]] = []
    for m in QUANTITY_RE.finditer(masked):
        start = m.start("num")
        if start > 0 and (masked[start - 1].isalnum() or masked[start - 1] in "_"):
            continue
        end = m.end("num")
        if end < len(masked) and masked[end:end + 1].isalpha() and not m.group("unit"):
            continue
        unit_raw = m.group("unit") or m.group("prefix")
        try:
            unit = canonical_unit(unit_raw) if unit_raw else None
        except ValueError:
            unit = None
        out.append((m.group(0).strip(), parse_number(m.group("num")), unit))
    return out


def _leaf_numbers(value: Any) -> Iterable[float]:
    if isinstance(value, bool):
        return
    if isinstance(value, (int, float)):
        yield float(value)
    elif isinstance(value, dict):
        for v in value.values():
            yield from _leaf_numbers(v)
    elif isinstance(value, list):
        for v in value:
            yield from _leaf_numbers(v)


class EvidenceIndex:
    def __init__(self, records: list[LedgerRecord]) -> None:
        self.records = [r for r in records if r.kind in PROVENANCE_KINDS]
        self.entries: list[tuple[float, str | None, LedgerRecord]] = []
        for r in self.records:
            text = r.body if isinstance(r.body, str) else json.dumps(r.body, ensure_ascii=False)
            for _raw, mag, unit in numbers_in(text):
                self.entries.append((mag, unit, r))
            if not isinstance(r.body, str):
                for mag in _leaf_numbers(r.body):
                    self.entries.append((mag, None, r))
            for tv in r.fields.values():
                if tv.magnitude is not None:
                    self.entries.append((tv.magnitude, tv.unit, r))

    @staticmethod
    def _close(a: float, ua: str | None, b: float, ub: str | None, tol: float) -> bool:
        if ua and ub:
            if not units_compatible(ua, ub):
                return False
            a, b = to_base(a, ua), to_base(b, ub)
        return math.isclose(a, b, rel_tol=tol, abs_tol=1e-9)

    def find(self, value: float, unit: str | None, tol: float, prefer: list[str]) -> LedgerRecord | None:
        matches = [r for mag, u, r in self.entries if self._close(value, unit, mag, u, tol)]
        if not matches:
            return None
        order = {rid: i for i, rid in enumerate(prefer)}

        def rank(r: LedgerRecord) -> tuple[int, int, int]:
            return (order.get(r.id, len(order)), 1 if r.kind in COMPUTE_KINDS else 0, r.seq)

        return min(matches, key=rank)


def _paragraph_text(paragraph: Any, refs: dict[str, str]) -> tuple[str, list[str]]:
    parts: list[str] = []
    cited: list[str] = []
    for run in paragraph.runs:
        if run.font.superscript:
            for m in REF_RE.finditer(run.text):
                cited += [refs[n.strip()] for n in m.group(1).split(",") if n.strip() in refs]
            parts.append(" ")
            continue
        parts.append(run.text)
    text = "".join(parts)
    for m in MARKER_RE.finditer(text):
        cited += [x.strip() for x in m.group(1).split(",")]
    return MARKER_RE.sub(" ", text), cited


REFERENCE_STYLE = "WB Reference"


def docx_references(doc: Any) -> dict[str, str]:
    refs: dict[str, str] = {}
    for p in doc.paragraphs:
        if p.style is not None and p.style.name == REFERENCE_STYLE:
            m = re.match(r"\[(\d+)\]\s+(R-[A-Za-z0-9]+-\d+)", p.text)
            if m:
                refs[m.group(1)] = m.group(2)
    return refs


def extract_docx(path: Path) -> list[Figure]:
    from docx import Document

    doc = Document(str(path))
    refs = docx_references(doc)
    figures: list[Figure] = []
    for i, p in enumerate(doc.paragraphs):
        if p.style is not None and p.style.name in {REFERENCE_STYLE, "WB Marking", "WB Signature", "WB Meta"}:
            continue
        text, cited = _paragraph_text(p, refs)
        for raw, mag, unit in numbers_in(text):
            figures.append(Figure(id=f"F{len(figures) + 1}", raw=raw, value=mag, unit=unit,
                                  location=f"paragraph {i + 1}", context=text.strip()[:300], cited=cited))
    for t_index, table in enumerate(doc.tables):
        for r_index, row in enumerate(table.rows):
            row_cited: list[str] = []
            texts = []
            for cell in row.cells:
                ctext = ""
                for p in cell.paragraphs:
                    ptext, pc = _paragraph_text(p, refs)
                    ctext += ptext + " "
                    row_cited += pc
                texts.append(ctext)
            for c_index, ctext in enumerate(texts):
                for raw, mag, unit in numbers_in(ctext):
                    figures.append(Figure(id=f"F{len(figures) + 1}", raw=raw, value=mag, unit=unit,
                                          location=f"table {t_index + 1} row {r_index + 1} col {c_index + 1}",
                                          context=" | ".join(x.strip() for x in texts)[:300], cited=row_cited))
    return figures


def extract_xlsx(path: Path) -> list[Figure]:
    from openpyxl import load_workbook

    wb = load_workbook(str(path), data_only=False)
    figures: list[Figure] = []
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for cell in row:
                v = cell.value
                comment = cell.comment.text if cell.comment else ""
                cited = re.findall(r"R-[A-Za-z0-9]+-\d+", comment)
                loc = f"{ws.title}!{cell.coordinate}"
                if isinstance(v, str) and v.startswith("="):
                    figures.append(Figure(id=f"F{len(figures) + 1}", raw=v, value=math.nan, location=loc,
                                          formula=v, cited=cited, context=comment))
                elif isinstance(v, (int, float)) and not isinstance(v, bool):
                    figures.append(Figure(id=f"F{len(figures) + 1}", raw=str(v), value=float(v), location=loc,
                                          cited=cited, context=comment))
    return figures


def resolve(figures: list[Figure], records: list[LedgerRecord], ledger: Ledger, tolerance: float,
            file: str) -> ProvenanceReport:
    index = EvidenceIndex(records)
    by_cell = {f.location: f for f in figures}
    out: list[Figure] = []
    for f in figures:
        if f.formula is not None:
            refs = re.findall(r"\$?([A-Z]{1,3})\$?(\d+)", f.formula)
            sheet = f.location.split("!")[0]
            chain_ids: list[str] = []
            for col, row in refs:
                src = by_cell.get(f"{sheet}!{col}{row}")
                if src is not None:
                    chain_ids += src.cited
            out.append(f.model_copy(update={"status": "derived", "chain": list(dict.fromkeys(chain_ids)),
                                            "note": "live formula"}))
            continue
        rec = index.find(f.value, f.unit, tolerance, f.cited)
        if rec is None:
            out.append(f.model_copy(update={"status": "unsourced"}))
            continue
        chain = [r.id for r in ledger.walk_inputs(rec.id)]
        status: Status = "derived" if rec.kind in COMPUTE_KINDS else "sourced"
        out.append(f.model_copy(update={"status": status, "record_id": rec.id, "chain": chain}))
    return ProvenanceReport(file=file, figures=out)


def check_file(path: Path, records: list[LedgerRecord], ledger: Ledger, tolerance: float) -> ProvenanceReport:
    suffix = path.suffix.lower()
    if suffix == ".docx":
        figures = extract_docx(path)
    elif suffix == ".xlsx":
        figures = extract_xlsx(path)
    else:
        figures = []
        text = path.read_text(encoding="utf-8", errors="replace") if suffix in {".md", ".txt"} else ""
        for i, line in enumerate(text.splitlines()):
            cited = [x.strip() for m in MARKER_RE.finditer(line) for x in m.group(1).split(",")]
            for raw, mag, unit in numbers_in(MARKER_RE.sub(" ", line)):
                figures.append(Figure(id=f"F{len(figures) + 1}", raw=raw, value=mag, unit=unit,
                                      location=f"line {i + 1}", context=line[:300], cited=cited))
    return resolve(figures, records, ledger, tolerance, path.name)
