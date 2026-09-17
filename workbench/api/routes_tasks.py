"""Tasks, approval gates, evidence, checks and draft review."""

from __future__ import annotations

import re
from pathlib import Path, PurePosixPath
from typing import Any, Literal

from fastapi import APIRouter, Depends, HTTPException
from fastapi.responses import FileResponse
from pydantic import BaseModel, Field

from workbench.agent.approvals import (
    ActionDecision,
    ChoiceDecision,
    DeliverableDecision,
    PlanDecision,
)
from workbench.agent.state import FigureDecision, TaskState
from workbench.api.deps import current_user, get_rt, require_role, task_for, workspace_for
from workbench.core.errors import NotFound, PolicyError
from workbench.core.labels import User
from workbench.core.normalise import parse_number
from workbench.planning.plan_schema import Plan
from workbench.runtime import Runtime

router = APIRouter(tags=["tasks"])


def record_view(rec: Any) -> dict[str, Any]:
    return {"id": rec.id, "seq": rec.seq, "kind": rec.kind, "trust": rec.trust, "summary": rec.summary,
            "label": rec.label.display(), "confidence": rec.confidence, "anchor": rec.anchor.model_dump()
            if rec.anchor else None, "source": rec.source(), "inputs": rec.inputs, "produced_by": rec.produced_by,
            "hash": rec.hash, "created_at": rec.created_at.isoformat()}


def task_view(rt: Runtime, state: TaskState, user: User, full: bool = True) -> dict[str, Any]:
    label = state.label or rt.ledger.high_water(state.id, state.label_floor)
    data: dict[str, Any] = {
        "id": state.id, "workspace": state.workspace, "user": state.user, "text": state.text,
        "attachments": state.attachments, "status": state.status, "status_note": state.status_note,
        "created_at": state.created_at, "updated_at": state.updated_at, "revision_no": state.revision_no,
        "model": state.model,
        "label": label.model_dump(mode="json"), "label_display": label.display(), "marking": label.marking(),
        "parent_id": state.parent_id, "children": state.children, "error": state.error,
        "followup_of": state.meta.get("followup_of"),
    }
    if not full:
        return data
    job = rt.jobs.for_task(state.id)
    queue = {j.id: j for j in rt.jobs.list(include_done=False)}
    pending = state.pending_gate()
    data.update({
        "route": state.route, "plan": state.plan.model_dump(mode="json") if state.plan else None,
        "plan_history": state.plan_history, "template_options": state.template_options,
        "trace": [t.model_dump(mode="json") for t in state.trace], "gates": [g.model_dump() for g in state.gates],
        "pending_gate": pending.model_dump() if pending else None, "escalations": state.escalations,
        "deliverables": [d.model_dump(mode="json") for d in state.deliverables], "checks": state.checks,
        "acknowledged": state.acknowledged, "result": state.result, "template_fallbacks": state.template_fallbacks,
        "job": (queue.get(job.id) or job).model_dump() if job else None, "meta": {
            k: v for k, v in state.meta.items() if k not in {"preread", "recalled"}},
        "can_approve": "approver" in user.roles, "is_owner": user.id == state.user,
        "draft_summary": rt.orchestrator.draft_summary(state) if state.deliverables else None,
        "blockers": rt.orchestrator.approval_blockers(state) if state.deliverables else [],
        "followups": [] if state.meta.get("followup_of") else _followup_views(rt, state, user),
    })
    return data


def _followup_views(rt: Runtime, state: TaskState, user: User) -> list[dict[str, Any]]:
    """The rest of the conversation: follow-ups the viewer may read, oldest first."""
    ws = rt.policy.workspace(state.workspace)
    out = []
    for f in rt.tasks.followups(state.id):
        label = f.label or rt.ledger.high_water(f.id, f.label_floor)
        if rt.policy.can_read(user, ws, label):
            out.append(task_view(rt, f, user))
    return out


class CreateTask(BaseModel):
    workspace: str
    text: str = Field(min_length=1, max_length=4000)
    attachments: list[str] = []
    meta: dict[str, Any] = {}


ALLOWED_META = {"report_date", "equipment_tag", "templates_disabled"}


@router.post("/tasks", status_code=201)
def create_task(body: CreateTask, user: User = Depends(current_user), rt: Runtime = Depends(get_rt)) -> dict[str, Any]:
    workspace_for(rt, user, body.workspace)
    meta = {k: v for k, v in body.meta.items() if k in ALLOWED_META}
    try:
        state = rt.orchestrator.create_task(body.workspace, user.id, body.text, body.attachments, meta)
    except (PolicyError, NotFound) as exc:
        raise HTTPException(403, str(exc)) from exc
    rt.jobs.submit(state.id)
    return task_view(rt, rt.tasks.get(state.id), user)


@router.get("/tasks")
def list_tasks(workspace: str | None = None, user: User = Depends(current_user),
               rt: Runtime = Depends(get_rt)) -> list[dict[str, Any]]:
    out = []
    for state in rt.tasks.list(workspace, limit=200):
        ws = rt.policy.workspaces.get(state.workspace)
        if ws is None or not rt.policy.can_access_workspace(user, ws):
            continue
        if state.label and not rt.policy.can_read(user, ws, state.label):
            continue
        out.append(task_view(rt, state, user, full=False))
    return out


@router.get("/tasks/{task_id}")
def get_task(task_id: str, user: User = Depends(current_user), rt: Runtime = Depends(get_rt)) -> dict[str, Any]:
    return task_view(rt, task_for(rt, user, task_id), user)


def _gate(state: TaskState, gate_id: str | None, kind: str) -> Any:
    gate = state.pending_gate()
    if gate is None or gate.kind != kind or (gate_id and gate.id != gate_id):
        raise HTTPException(409, f"task {state.id} is not waiting for a {kind.replace('_', ' ')} decision")
    return gate


def _decide(rt: Runtime, gate_id: str, decision: Any) -> None:
    decide = getattr(rt.approvals, "decide", None)
    if decide is None:
        raise HTTPException(409, "this server runs with automatic approvals")
    try:
        decide(gate_id, decision)
    except NotFound as exc:
        raise HTTPException(409, str(exc)) from exc


class PlanDecisionBody(BaseModel):
    decision: Literal["approve", "edit", "reject"]
    plan: dict[str, Any] | None = None
    note: str | None = None


@router.post("/tasks/{task_id}/plan/decision")
def plan_decision(task_id: str, body: PlanDecisionBody, user: User = Depends(current_user),
                  rt: Runtime = Depends(get_rt)) -> dict[str, Any]:
    state = task_for(rt, user, task_id)
    if user.id != state.user and "approver" not in user.roles:
        raise HTTPException(403, "only the task owner or an approver can decide the plan")
    gate = _gate(state, None, "plan")
    plan = None
    if body.decision == "approve" and state.plan is not None and not state.plan.valid:
        raise HTTPException(409, "the plan has problems: " + "; ".join(state.plan.errors))
    if body.decision == "edit":
        if body.plan is None:
            raise HTTPException(422, "an edit needs the edited plan")
        try:
            plan = Plan.from_typed(body.plan)
        except (ValueError, KeyError) as exc:
            raise HTTPException(422, f"invalid plan: {exc}") from exc
    _decide(rt, gate.id, PlanDecision(body.decision, plan=plan, by=user.id, note=body.note))
    return {"ok": True, "gate": gate.id}


class ChoiceBody(BaseModel):
    template: str


@router.post("/tasks/{task_id}/template/decision")
def template_decision(task_id: str, body: ChoiceBody, user: User = Depends(current_user),
                      rt: Runtime = Depends(get_rt)) -> dict[str, Any]:
    state = task_for(rt, user, task_id)
    gate = _gate(state, None, "template_choice")
    if body.template not in state.template_options:
        raise HTTPException(422, f"choose one of {state.template_options}")
    _decide(rt, gate.id, ChoiceDecision(body.template, by=user.id))
    return {"ok": True}


class ActionBody(BaseModel):
    approve: bool
    note: str | None = None


@router.post("/tasks/{task_id}/actions/{aid}/decision")
def action_decision(task_id: str, aid: str, body: ActionBody, user: User = Depends(current_user),
                    rt: Runtime = Depends(get_rt)) -> dict[str, Any]:
    state = task_for(rt, user, task_id)
    if user.id != state.user and "approver" not in user.roles:
        raise HTTPException(403, "only the task owner or an approver can approve actions")
    gate = _gate(state, aid, "action")
    _decide(rt, gate.id, ActionDecision(body.approve, by=user.id, note=body.note))
    return {"ok": True}


@router.get("/tasks/{task_id}/ledger")
def ledger(task_id: str, user: User = Depends(current_user), rt: Runtime = Depends(get_rt)) -> list[dict[str, Any]]:
    state = task_for(rt, user, task_id)
    ids = [state.id, *state.children]
    return [record_view(r) for tid in ids for r in rt.ledger.for_task(tid)]


@router.get("/records/{rid}")
def record(rid: str, user: User = Depends(current_user), rt: Runtime = Depends(get_rt)) -> dict[str, Any]:
    try:
        rec = rt.ledger.get(rid)
    except NotFound as exc:
        raise HTTPException(404, str(exc)) from exc
    state = task_for(rt, user, rec.task_id)
    if not rt.policy.can_read(user, rt.policy.workspace(state.workspace), rec.label):
        raise HTTPException(403, f"record {rid} is {rec.label.display()}, above your clearance")
    chain = [record_view(r) for r in rt.ledger.walk_inputs(rid)]
    return {**record_view(rec), "body": rec.body, "fields": {k: v.model_dump() for k, v in rec.fields.items()},
            "chain": chain}


@router.get("/tasks/{task_id}/evidence/{name}")
def evidence(task_id: str, name: str, user: User = Depends(current_user), rt: Runtime = Depends(get_rt)) -> FileResponse:
    task_for(rt, user, task_id)
    if not re.fullmatch(r"[A-Za-z0-9_.-]+\.png", name):
        raise HTTPException(404, "not found")
    path = rt.evidence_dir(task_id) / name
    if not path.is_file():
        raise HTTPException(404, "not found")
    return FileResponse(path, media_type="image/png")


@router.get("/tasks/{task_id}/checks")
def checks(task_id: str, user: User = Depends(current_user), rt: Runtime = Depends(get_rt)) -> dict[str, Any]:
    state = task_for(rt, user, task_id)
    records = rt.ledger.for_task(task_id)
    rows = [r for r in records if r.kind == "check_result"]
    out = []
    for rec in rows:
        body = dict(rec.body) if isinstance(rec.body, dict) else {}
        body["record_id"] = rec.id
        body["acknowledged"] = rec.id in state.acknowledged
        extra = body.get("extra") or {}
        left = rt.ledger.maybe(body.get("left_record") or "")
        if left is not None and left.kind == "vlm_read" and isinstance(left.body, dict):
            extra = {**extra, "crop": left.body.get("crop"), "zoomed_crop": left.body.get("zoomed_crop")}
        body["extra"] = extra
        out.append(body)
    uncertain = [record_view(r) | {"body": r.body} for r in records if r.kind == "vlm_read"]
    return {"checks": out, "dual_read": uncertain}


@router.post("/tasks/{task_id}/checks/{rid}/acknowledge")
def acknowledge(task_id: str, rid: str, user: User = Depends(current_user), rt: Runtime = Depends(get_rt)) -> dict[str, Any]:
    state = task_for(rt, user, task_id)
    require_role(user, "approver")
    if rid not in [c.get("record_id") for c in state.checks]:
        raise HTTPException(404, f"{rid} is not a check of this task")
    if rid not in state.acknowledged:
        state.acknowledged.append(rid)
        rt.tasks.save(state)
        rt.audit.append({"type": "check.acknowledged", "task": task_id, "check": rid, "by": user.id})
    return {"acknowledged": state.acknowledged}


def _deliverable_preview(rt: Runtime, state: TaskState, d: Any) -> dict[str, Any]:
    path = rt.files.open_path(d.file_id)
    suffix = path.suffix.lower()
    preview: dict[str, Any] = {"type": suffix.lstrip(".")}
    if suffix == ".docx":
        from workbench.tools.render_docx import docx_to_blocks

        preview["blocks"] = docx_to_blocks(path)
    elif suffix == ".xlsx":
        from openpyxl import load_workbook

        wb = load_workbook(str(path))
        sheets = []
        for ws in wb.worksheets:
            rows = []
            for row in ws.iter_rows(max_row=min(ws.max_row, 60), max_col=min(ws.max_column, 16)):
                rows.append([{"v": c.value if c.value is None or isinstance(c.value, (int, float, str)) else str(c.value),
                              "ref": c.coordinate, "comment": c.comment.text if c.comment else None} for c in row])
            sheets.append({"name": ws.title, "rows": rows})
        preview["sheets"] = sheets
    elif suffix in {".py", ".md", ".txt", ".csv", ".json"}:
        preview["text"] = path.read_text(encoding="utf-8", errors="replace")[:20000]
    elif suffix == ".svg":
        preview["svg"] = path.read_text(encoding="utf-8", errors="replace")[:200000]
    elif suffix == ".pptx":
        from pptx import Presentation

        preview["slides"] = [[sh.text_frame.text for sh in s.shapes if sh.has_text_frame and sh.name != "WB Marking"]
                             for s in Presentation(str(path)).slides]
    return {**d.model_dump(mode="json"), "preview": preview, "label_display": d.label.display(),
            "marking": d.label.marking(), "name": PurePosixPath(d.relpath).name}


@router.get("/tasks/{task_id}/draft")
def draft(task_id: str, user: User = Depends(current_user), rt: Runtime = Depends(get_rt)) -> dict[str, Any]:
    state = task_for(rt, user, task_id)
    if not state.deliverables:
        raise HTTPException(404, "this task has no draft deliverables yet")
    return {"task": task_view(rt, state, user, full=False),
            "deliverables": [_deliverable_preview(rt, state, d) for d in state.deliverables],
            "checks": state.checks, "acknowledged": state.acknowledged,
            "figure_decisions": [f.model_dump() for f in state.figure_decisions],
            "summary": rt.orchestrator.draft_summary(state), "blockers": rt.orchestrator.approval_blockers(state),
            "pending_gate": state.pending_gate().model_dump() if state.pending_gate() else None,
            "can_approve": "approver" in user.roles}


class FigureBody(BaseModel):
    action: Literal["correct", "link", "confirm"]
    record_id: str | None = None
    value: str | None = None
    note: str | None = None


def _correct_in_file(path: Path, location: str, raw: str, value: str) -> bool:
    if path.suffix.lower() != ".docx":
        return False
    from docx import Document

    doc = Document(str(path))
    targets = []
    m = re.match(r"paragraph (\d+)", location)
    t = re.match(r"table (\d+) row (\d+) col (\d+)", location)
    if m:
        targets = [doc.paragraphs[int(m.group(1)) - 1]]
    elif t:
        cell = doc.tables[int(t.group(1)) - 1].rows[int(t.group(2)) - 1].cells[int(t.group(3)) - 1]
        targets = list(cell.paragraphs)
    for p in targets:
        for r in p.runs:
            if raw in r.text:
                r.text = r.text.replace(raw, value, 1)
                doc.save(str(path))
                return True
    return False


@router.post("/tasks/{task_id}/draft/figures/{fid}")
def figure_decision(task_id: str, fid: str, body: FigureBody, user: User = Depends(current_user),
                    rt: Runtime = Depends(get_rt)) -> dict[str, Any]:
    state = task_for(rt, user, task_id)
    require_role(user, "approver")
    file_id, _, figure_id = fid.partition(":")
    deliverable = next((d for d in state.deliverables if d.file_id == file_id), None)
    figure = next((f for f in (deliverable.provenance.get("figures", []) if deliverable else [])
                   if f["id"] == figure_id), None)
    if deliverable is None or figure is None:
        raise HTTPException(404, f"figure {fid} not found")
    if body.action == "link":
        rec = rt.ledger.maybe(body.record_id or "")
        if rec is None or rec.task_id not in {task_id, *state.children}:
            raise HTTPException(422, "link needs a record of this task")
        if f"{figure['value']:g}" not in rec.body_text() and figure["raw"].split()[0] not in rec.body_text():
            raise HTTPException(422, f"{rec.id} does not contain {figure['raw']}")
    if body.action == "correct":
        if not body.value:
            raise HTTPException(422, "a correction needs the corrected value")
        try:
            parse_number(body.value.split()[0])
        except ValueError as exc:
            raise HTTPException(422, "the corrected value must start with a number") from exc
        if not _correct_in_file(rt.files.open_path(file_id), figure["location"], figure["raw"], body.value):
            raise HTTPException(422, "the figure could not be located in the document")
        rt.files.register(deliverable_ws(rt, file_id), deliverable.relpath, deliverable.label, task_id)
    decision = FigureDecision(figure=fid, action=body.action, record_id=body.record_id, value=body.value,
                              note=body.note, by=user.id, at=rt.clock.now().isoformat())
    state.figure_decisions = [f for f in state.figure_decisions if f.figure != fid] + [decision]
    rt.tasks.save(state)
    rt.audit.append({"type": "figure.decision", "task": task_id, "figure": fid, "raw": figure["raw"],
                     "action": body.action, "record": body.record_id, "value": body.value, "note": body.note,
                     "by": user.id})
    if body.action == "correct":
        rt.orchestrator.review(state)
        fresh = rt.tasks.get(task_id)
        fresh.figure_decisions = state.figure_decisions
        rt.tasks.save(fresh)
    return {"ok": True, "summary": rt.orchestrator.draft_summary(rt.tasks.get(task_id))}


def deliverable_ws(rt: Runtime, file_id: str) -> str:
    return rt.files.get(file_id).workspace


class DraftDecisionBody(BaseModel):
    decision: Literal["approve", "reject"]
    acknowledged: list[str] = []
    note: str | None = None


@router.post("/tasks/{task_id}/draft/decision")
def draft_decision(task_id: str, body: DraftDecisionBody, user: User = Depends(current_user),
                   rt: Runtime = Depends(get_rt)) -> dict[str, Any]:
    state = task_for(rt, user, task_id)
    require_role(user, "approver")
    gate = _gate(state, None, "deliverable")
    if body.decision == "approve":
        valid = {c.get("record_id") for c in state.checks}
        for rid in body.acknowledged:
            if rid in valid and rid not in state.acknowledged:
                state.acknowledged.append(rid)
                rt.audit.append({"type": "check.acknowledged", "task": task_id, "check": rid, "by": user.id})
        rt.tasks.save(state)
        blockers = rt.orchestrator.approval_blockers(state)
        if blockers:
            raise HTTPException(409, "cannot approve yet: " + "; ".join(blockers))
    elif not body.note:
        raise HTTPException(422, "a rejection needs a note for the revision")
    _decide(rt, gate.id, DeliverableDecision(body.decision, by=user.id, note=body.note,
                                             acknowledged=list(state.acknowledged)))
    return {"ok": True}


class FollowUp(BaseModel):
    text: str = Field(min_length=3, max_length=2000)


@router.post("/tasks/{task_id}/followup", status_code=201)
def followup(task_id: str, body: FollowUp, user: User = Depends(current_user),
             rt: Runtime = Depends(get_rt)) -> dict[str, Any]:
    parent = task_for(rt, user, task_id)
    root = parent
    while root.meta.get("followup_of"):
        root = task_for(rt, user, str(root.meta["followup_of"]))
    floor = rt.ledger.high_water(root.id, root.label_floor)
    for earlier in rt.tasks.followups(root.id):
        floor = floor.join(rt.ledger.high_water(earlier.id, earlier.label_floor))
    try:
        # Follow-ups join the conversation's root and start at its classification.
        child = rt.orchestrator.create_task(root.workspace, user.id, body.text, root.attachments,
                                            meta={"followup_of": root.id}, parent_id=root.id, label_floor=floor)
    except PolicyError as exc:
        raise HTTPException(403, str(exc)) from exc
    rt.jobs.submit(child.id)
    return task_view(rt, rt.tasks.get(child.id), user)


@router.get("/routing/{task_id}")
def routing(task_id: str, user: User = Depends(current_user), rt: Runtime = Depends(get_rt)) -> dict[str, Any]:
    state = task_for(rt, user, task_id)
    if not state.route:
        raise HTTPException(404, "the task has not been routed yet")
    return state.route
