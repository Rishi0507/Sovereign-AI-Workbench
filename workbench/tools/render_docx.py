"""``make_docx``: approval notes, recommendation notes, summaries, answers and calculation sheets.

Built from ``org_templates/approval_note.docx``. The inherited label is stamped in the header,
the footer, a marking line at the top and the core properties. Citation markers ``[R-...]``
become numbered superscript references listed under "Evidence references".
"""

from __future__ import annotations

import io
import re
from pathlib import Path
from typing import Any

from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt, RGBColor

from workbench.checks.provenance import MARKER_RE, REFERENCE_STYLE
from workbench.core.errors import ToolError
from workbench.core.labels import Label
from workbench.core.ledger import Ledger
from workbench.tools.registry import ToolContext, ToolResult, ToolSpec, obj
from workbench.tools.render_common import MARKING_PLACEHOLDER, label_properties, marking_text

SCHEMA = obj({
    "template": {"type": "string", "minLength": 1},
    "data": {"type": "object"},
    "name": {"type": "string", "pattern": r"^[A-Za-z0-9_.-]+\.docx$"},
}, ["template", "data"])

KINDS = {"approval_note": "approval-note.docx", "recommendation_note": "recommendation-note.docx",
         "summary": "summary.docx", "answer": "answer.docx", "calc_sheet": "calc-sheet.docx"}
STATUS_TEXT = {"pass": "Pass", "mismatch": "Mismatch", "not_found": "Not found", "not_checked": "Not checked"}


def template_kind(name: str) -> str:
    stem = Path(name).stem.replace("-", "_")
    for kind in KINDS:
        if stem == kind or stem.endswith(kind):
            return kind
    raise ToolError(f"unknown document template {name!r}; known: {', '.join(KINDS)}")


class RefBook:
    def __init__(self, ledger: Ledger) -> None:
        self.ledger = ledger
        self.numbers: dict[str, int] = {}

    def number(self, rid: str) -> int:
        if rid not in self.numbers:
            self.numbers[rid] = len(self.numbers) + 1
        return self.numbers[rid]

    def cite(self, rid: str) -> str:
        rec = self.ledger.maybe(rid)
        if rec is None:
            return "record not found"
        return f"{rec.kind}, {rec.source()}"


def _styles(doc: Any) -> None:
    styles = doc.styles
    for name, size, bold, color in (("WB Marking", 9, True, RGBColor(0x9B, 0x1C, 0x1C)),
                                    (REFERENCE_STYLE, 8.5, False, RGBColor(0x44, 0x4B, 0x55)),
                                    ("WB Signature", 10, False, RGBColor(0x1F, 0x2A, 0x2E)),
                                    ("WB Meta", 9.5, False, RGBColor(0x44, 0x4B, 0x55))):
        if name not in [s.name for s in styles]:
            st = styles.add_style(name, WD_STYLE_TYPE.PARAGRAPH)
            st.base_style = styles["Normal"]
            st.font.size = Pt(size)
            st.font.bold = bold
            st.font.color.rgb = color


def add_cited(paragraph: Any, text: str, refs: RefBook) -> None:
    pos = 0
    for m in MARKER_RE.finditer(text):
        before = text[pos:m.start()].rstrip()
        if before:
            paragraph.add_run(before)
        ids = [x.strip() for x in m.group(1).split(",")]
        run = paragraph.add_run("[" + ",".join(str(refs.number(i)) for i in ids) + "]")
        run.font.superscript = True
        pos = m.end()
    tail = text[pos:]
    if tail:
        paragraph.add_run(tail)


def new_document(template: Path, label: Label, title: str, subject: str) -> Any:
    doc = Document(str(template)) if template.is_file() else Document()
    _styles(doc)
    marking = marking_text(label)
    for section in doc.sections:
        for part in (section.header, section.footer):
            if not part.paragraphs[0].text.strip():
                part.paragraphs[0].text = MARKING_PLACEHOLDER
            for p in part.paragraphs:
                if MARKING_PLACEHOLDER in p.text:
                    p.text = p.text.replace(MARKING_PLACEHOLDER, marking)
                    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    body = doc.element.body
    for p in list(doc.paragraphs):
        if "{{BODY}}" in p.text:
            body.remove(p._element)
    mp = doc.add_paragraph(marking, style="WB Marking")
    mp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    doc.add_heading(title, level=1)
    cp = doc.core_properties
    props = label_properties(label)
    cp.title, cp.subject = title, subject
    cp.keywords, cp.category, cp.comments = props["keywords"], props["category"], props["comments"]
    cp.author = "Sovereign AI Workbench"
    return doc


def finish_document(doc: Any, refs: RefBook, signature: bool = True) -> bytes:
    if signature:
        doc.add_heading("Approval", level=2)
        for line in ("Prepared by: ______________________   Date: __________",
                     "Reviewed by: ______________________   Date: __________",
                     "Approved by: ______________________   Date: __________"):
            doc.add_paragraph(line, style="WB Signature")
    if refs.numbers:
        doc.add_heading("Evidence references", level=2)
        for rid, n in sorted(refs.numbers.items(), key=lambda kv: kv[1]):
            doc.add_paragraph(f"[{n}] {rid}  {refs.cite(rid)}", style=REFERENCE_STYLE)
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def render_note(data: dict[str, Any], label: Label, template: Path, refs: RefBook, reference: str) -> bytes:
    doc = new_document(template, label, str(data.get("title") or "Note"), str(data.get("subject") or ""))
    doc.add_paragraph(f"Reference: {data.get('reference') or reference}", style="WB Meta")
    doc.add_paragraph(f"Subject: {data.get('subject', '')}", style="WB Meta")
    doc.add_paragraph(f"Classification: {label.display()}", style="WB Meta")
    sections = [("Summary", data.get("summary") or []), ("Findings", data.get("findings") or [])]
    n = 0
    for heading, items in sections:
        if not items:
            continue
        n += 1
        doc.add_heading(f"{n}. {heading}", level=2)
        for item in items:
            add_cited(doc.add_paragraph(style="List Bullet"), str(item.get("text", "")), refs)
    checks = data.get("consistency_findings") or []
    n += 1
    doc.add_heading(f"{n}. Consistency findings", level=2)
    if checks:
        table = doc.add_table(rows=1, cols=3)
        table.style = "Light Grid Accent 1" if "Light Grid Accent 1" in [s.name for s in doc.styles] else table.style
        hdr = table.rows[0].cells
        hdr[0].text, hdr[1].text, hdr[2].text = "Check", "Result", "Detail"
        for c in checks:
            row = table.add_row().cells
            row[0].text = str(c.get("rule", "")).replace("_", " ")
            row[1].text = STATUS_TEXT.get(str(c.get("status")), str(c.get("status")))
            add_cited(row[2].paragraphs[0], str(c.get("text", "")), refs)
    else:
        doc.add_paragraph("No consistency checks apply to this deliverable.")
    if data.get("recommendation"):
        n += 1
        doc.add_heading(f"{n}. Recommendation", level=2)
        for item in data["recommendation"]:
            add_cited(doc.add_paragraph(), str(item.get("text", "")), refs)
    if data.get("incomplete"):
        n += 1
        doc.add_heading(f"{n}. Incomplete steps", level=2)
        for item in data["incomplete"]:
            doc.add_paragraph(str(item), style="List Bullet")
    return finish_document(doc, refs)


def render_summary(data: dict[str, Any], label: Label, template: Path, refs: RefBook) -> bytes:
    doc = new_document(template, label, str(data.get("title") or "Summary"), "Document summary")
    if data.get("overall"):
        doc.add_heading("Key points", level=2)
        for pt in data["overall"]:
            add_cited(doc.add_paragraph(style="List Bullet"), str(pt.get("text", "")), refs)
    for sec in data.get("sections") or []:
        doc.add_heading(str(sec.get("heading", "")), level=2)
        for pt in sec.get("points") or []:
            add_cited(doc.add_paragraph(style="List Bullet"), str(pt.get("text", "")), refs)
    return finish_document(doc, refs, signature=False)


def render_answer(data: dict[str, Any], label: Label, template: Path, refs: RefBook) -> bytes:
    doc = new_document(template, label, str(data.get("title") or "Answer"), str(data.get("question") or ""))
    if data.get("question"):
        doc.add_paragraph(f"Question: {data['question']}", style="WB Meta")
    for pt in data.get("answer") or []:
        add_cited(doc.add_paragraph(), str(pt.get("text", "")), refs)
    return finish_document(doc, refs, signature=False)


def render_calc(data: dict[str, Any], label: Label, template: Path, refs: RefBook) -> bytes:
    calc = data.get("calc") or data
    doc = new_document(template, label, str(calc.get("title") or "Calculation sheet"), str(calc.get("expression")))
    doc.add_heading("Inputs", level=2)
    table = doc.add_table(rows=1, cols=4)
    for cell, text in zip(table.rows[0].cells, ("Symbol", "Description", "Value", "Source"), strict=True):
        cell.text = text
    for sym, v in (calc.get("variables") or {}).items():
        row = table.add_row().cells
        row[0].text, row[1].text = sym, str(v.get("description", ""))
        row[2].text = f"{v.get('value')} {v.get('unit')}"
        if v.get("record"):
            add_cited(row[3].paragraphs[0], f"[{v['record']}]", refs)
    doc.add_heading("Steps", level=2)
    record = data.get("record")
    for st in calc.get("steps") or []:
        text = " ".join(x for x in (st.get("formula"), st.get("substitution"), st.get("result")) if x)
        line = f"Step {st.get('step')}: {st.get('label')}: {text}"
        if record and st.get("result"):
            line += f" [{record}]"
        add_cited(doc.add_paragraph(), line, refs)
    return finish_document(doc, refs)


def make_docx(args: dict[str, Any], ctx: ToolContext) -> ToolResult:
    kind = template_kind(str(args["template"]))
    data = dict(args["data"])
    if isinstance(data.get("value"), dict) and "records" in data:
        data = {**data["value"], **{k: v for k, v in data.items() if k not in {"value", "records"}}}
    label = ctx.label()
    template = ctx.rt.settings.root / "org_templates" / "approval_note.docx"
    refs = RefBook(ctx.rt.ledger)
    reference = f"WB/{ctx.workspace}/{ctx.task_id}"
    if kind in {"approval_note", "recommendation_note"}:
        content = render_note(data, label, template, refs, reference)
    elif kind == "summary":
        content = render_summary(data, label, template, refs)
    elif kind == "answer":
        content = render_answer(data, label, template, refs)
    else:
        content = render_calc(data, label, template, refs)
    name = str(args.get("name") or KINDS[kind])
    rec = ctx.rt.files.write_draft(ctx.workspace, name, content, label, ctx.task_id,
                                   overwrite_ok=bool(ctx.meta.get("action_approved", True)))
    return ToolResult(ok=True, summary=f"rendered drafts/{name} ({label.display()}, {len(refs.numbers)} reference(s))",
                      files=[rec.id], body={"file": rec.relpath, "file_id": rec.id, "label": label.display(),
                                            "references": refs.numbers, "kind": kind})


def docx_to_blocks(path: Path) -> list[dict[str, Any]]:
    """Read a rendered docx back into simple blocks for the browser preview."""
    doc = Document(str(path))
    refs: dict[str, str] = {}
    for p in doc.paragraphs:
        if p.style is not None and p.style.name == REFERENCE_STYLE:
            m = re.match(r"\[(\d+)\]\s+(R-[A-Za-z0-9]+-\d+)", p.text)
            if m:
                refs[m.group(1)] = m.group(2)
    blocks: list[dict[str, Any]] = []
    header = doc.sections[0].header.paragraphs[0].text if doc.sections else ""
    blocks.append({"type": "marking", "text": header})
    body = doc.element.body
    tables = iter(doc.tables)
    paragraphs = iter(doc.paragraphs)
    p_index = 0
    for child in body.iterchildren():
        tag = child.tag.rsplit("}", 1)[-1]
        if tag == "p":
            p = next(paragraphs)
            p_index += 1
            runs = [{"text": r.text, "sup": bool(r.font.superscript),
                     "refs": [refs.get(x.strip()) for x in re.findall(r"\d+", r.text)] if r.font.superscript else []}
                    for r in p.runs]
            style = p.style.name if p.style is not None else "Normal"
            if not p.text.strip():
                continue
            blocks.append({"type": "p", "style": style, "runs": runs, "index": p_index})
        elif tag == "tbl":
            t = next(tables)
            rows = []
            for row in t.rows:
                cells = []
                for cell in row.cells:
                    cell_runs = []
                    for p in cell.paragraphs:
                        for r in p.runs:
                            cell_runs.append({"text": r.text, "sup": bool(r.font.superscript),
                                              "refs": [refs.get(x) for x in re.findall(r"\d+", r.text)]
                                              if r.font.superscript else []})
                    cells.append(cell_runs)
                rows.append(cells)
            blocks.append({"type": "table", "rows": rows})
    return blocks


SPECS = [ToolSpec(name="make_docx", description="Render a Word deliverable from the org template into drafts/.",
                  input_schema=SCHEMA, handler=make_docx, side_effect=True, output_type="file",
                  budget_key="make_docx")]
