"""``make_pptx``: a briefing deck, one idea per slide, with the marking on every slide.

The deck is drawn rather than typed into a layout: a title bar, a rule under it, the points,
and a footer carrying the deck title, the classification and the slide number. Anything the
outline provides (a lead sentence, a source note) is placed where a reader expects it, and the
speaker notes carry the source so a slide can be traced back like any other deliverable.
"""

from __future__ import annotations

import io
from typing import Any

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Emu, Pt

from workbench.tools.registry import ToolContext, ToolResult, ToolSpec, obj
from workbench.tools.render_common import label_properties, marking_text

SCHEMA = obj({
    "data": {"type": "object"},
    "name": {"type": "string", "pattern": r"^[A-Za-z0-9_.-]+\.pptx$"},
}, ["data"])

INK = RGBColor(0x1F, 0x24, 0x2B)
MUTED = RGBColor(0x6B, 0x71, 0x7A)
ACCENT = RGBColor(0x1D, 0x4E, 0xD8)
MARKING_RED = RGBColor(0x9B, 0x1C, 0x1C)
FACE = "Calibri"


def _text(slide: Any, left: float, top: float, width: float, height: float, name: str = "") -> Any:
    box = slide.shapes.add_textbox(Emu(int(left)), Emu(int(top)), Emu(int(width)), Emu(int(height)))
    if name:
        box.name = name
    frame = box.text_frame
    frame.word_wrap = True
    return frame


def _style(run: Any, size: int, bold: bool = False, colour: RGBColor = INK) -> None:
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = colour
    run.font.name = FACE


def _rule(slide: Any, left: float, top: float, width: float, height: float, colour: RGBColor) -> None:
    from pptx.enum.shapes import MSO_SHAPE

    bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Emu(int(left)), Emu(int(top)), Emu(int(width)), Emu(int(height)))
    bar.fill.solid()
    bar.fill.fore_color.rgb = colour
    bar.line.fill.background()
    bar.shadow.inherit = False


def _footer(slide: Any, prs: Any, marking: str, deck_title: str, number: int | None) -> None:
    width, height = prs.slide_width, prs.slide_height
    margin = int(width * 0.06)
    # The marking box holds the marking and nothing else, so it can be read back as the classification.
    frame = _text(slide, margin, int(height * 0.90), int(width * 0.34), int(height * 0.06), name="WB Marking")
    run = frame.paragraphs[0].add_run()
    run.text = marking
    _style(run, 9, True, MARKING_RED)
    if deck_title:
        tail = _text(slide, int(width * 0.40), int(height * 0.90), int(width * 0.44), int(height * 0.06),
                     name="WB Footer")
        tail.paragraphs[0].alignment = PP_ALIGN.RIGHT
        run = tail.paragraphs[0].add_run()
        run.text = deck_title
        _style(run, 9, False, MUTED)
    if number is not None:
        right = _text(slide, int(width * 0.86), int(height * 0.90), int(width * 0.08), int(height * 0.06),
                      name="WB Number")
        right.paragraphs[0].alignment = PP_ALIGN.RIGHT
        run = right.paragraphs[0].add_run()
        run.text = str(number)
        _style(run, 9, False, MUTED)


def _notes(slide: Any, text: str) -> None:
    if text:
        slide.notes_slide.notes_text_frame.text = text


def _title_slide(prs: Any, data: dict[str, Any], label_line: str, marking: str) -> None:
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    width, height = prs.slide_width, prs.slide_height
    margin = int(width * 0.08)
    _rule(slide, margin, int(height * 0.30), int(width * 0.10), Emu(60000), ACCENT)
    frame = _text(slide, margin, int(height * 0.34), int(width * 0.84), int(height * 0.22), name="WB Title")
    run = frame.paragraphs[0].add_run()
    run.text = str(data.get("title", "Briefing"))
    _style(run, 36, True)
    subtitle = str(data.get("subtitle") or label_line)
    if subtitle:
        sub = _text(slide, margin, int(height * 0.57), int(width * 0.84), int(height * 0.10), name="WB Subtitle")
        run = sub.paragraphs[0].add_run()
        run.text = subtitle
        _style(run, 14, False, MUTED)
    _footer(slide, prs, marking, "", None)


def _content_slide(prs: Any, spec: dict[str, Any], marking: str, deck_title: str, number: int) -> None:
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    width, height = prs.slide_width, prs.slide_height
    margin = int(width * 0.08)
    frame = _text(slide, margin, int(height * 0.12), int(width * 0.84), int(height * 0.12), name="WB Title")
    run = frame.paragraphs[0].add_run()
    run.text = str(spec.get("title", ""))
    _style(run, 26, True)
    _rule(slide, margin, int(height * 0.255), int(width * 0.06), Emu(48000), ACCENT)

    lead = str(spec.get("lead") or "")
    top = int(height * 0.32)
    if lead:
        lead_frame = _text(slide, margin, top, int(width * 0.84), int(height * 0.10), name="WB Lead")
        run = lead_frame.paragraphs[0].add_run()
        run.text = lead
        _style(run, 15, False, MUTED)
        top = int(height * 0.43)

    bullets = [str(b).strip() for b in spec.get("bullets") or [] if str(b).strip()][:6]
    if bullets:
        body = _text(slide, margin, top, int(width * 0.84), int(height * 0.44), name="WB Body")
        body.vertical_anchor = MSO_ANCHOR.TOP
        for i, text in enumerate(bullets):
            para = body.paragraphs[0] if i == 0 else body.add_paragraph()
            para.space_after = Pt(10)
            dot = para.add_run()
            dot.text = "•   "
            _style(dot, 16, False, ACCENT)
            run = para.add_run()
            run.text = text
            _style(run, 16)
    _footer(slide, prs, marking, deck_title, number)
    _notes(slide, str(spec.get("note") or ""))


def make_pptx(args: dict[str, Any], ctx: ToolContext) -> ToolResult:
    data = dict(args["data"])
    if isinstance(data.get("value"), dict):
        data = data["value"]
    label = ctx.label()
    marking = marking_text(label)
    template = ctx.rt.settings.root / "org_templates" / "deck.pptx"
    prs = Presentation(str(template)) if template.is_file() else Presentation()
    for sld_id in list(prs.slides._sldIdLst):
        prs.part.drop_rel(sld_id.rId)
        prs.slides._sldIdLst.remove(sld_id)

    deck_title = str(data.get("title", "Briefing"))
    _title_slide(prs, data, label.display(), marking)
    for number, spec in enumerate(data.get("slides") or [], start=2):
        _content_slide(prs, dict(spec), marking, deck_title, number)

    cp = prs.core_properties
    props = label_properties(label)
    cp.title = deck_title
    cp.keywords, cp.category, cp.comments = props["keywords"], props["category"], props["comments"]
    buf = io.BytesIO()
    prs.save(buf)
    name = str(args.get("name") or "deck.pptx")
    rec = ctx.rt.files.write_draft(ctx.workspace, name, buf.getvalue(), label, ctx.task_id,
                                   overwrite_ok=bool(ctx.meta.get("action_approved", True)))
    return ToolResult(ok=True, summary=f"rendered drafts/{name} with {len(prs.slides)} slide(s)",
                      files=[rec.id], body={"file": rec.relpath, "file_id": rec.id, "label": label.display()})


SPECS = [ToolSpec(name="make_pptx", description="Render a slide deck from the org slide master into drafts/.",
                  input_schema=SCHEMA, handler=make_pptx, side_effect=True, output_type="file",
                  budget_key="make_pptx")]
