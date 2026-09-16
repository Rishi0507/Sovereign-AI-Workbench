"""``make_xlsx``: spreadsheets with live formulas for derived cells and sourced input cells.

Input cells carry a comment ``source: R-... (doc p.N)``; derived cells keep formulas, so the
derivation survives outside the workbench.
"""

from __future__ import annotations

import io
from typing import Any

from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from workbench.core.errors import ToolError
from workbench.core.labels import Label
from workbench.core.ledger import Ledger
from workbench.core.normalise import find_quantities, ureg
from workbench.tools.calculate import _unit, to_excel
from workbench.tools.registry import ToolContext, ToolResult, ToolSpec, obj
from workbench.tools.render_common import label_properties, marking_text

SCHEMA = obj({
    "kind": {"enum": ["calc_sheet", "offer_comparison", "findings", "table"]},
    "data": {"type": "object"},
    "name": {"type": "string", "pattern": r"^[A-Za-z0-9_.-]+\.xlsx$"},
}, ["kind", "data"])

HEAD_FONT = Font(bold=True, color="FFFFFF")
HEAD_FILL = PatternFill("solid", fgColor="1F3A3D")
MARK_FONT = Font(bold=True, color="9B1C1C")
INPUT_FILL = PatternFill("solid", fgColor="EEF6F4")


def _unwrap(value: Any) -> Any:
    if isinstance(value, dict) and "value" in value and "records" in value:
        return value["value"]
    return value


class Sheeter:
    def __init__(self, ledger: Ledger, label: Label) -> None:
        self.ledger = ledger
        self.label = label
        self.wb = Workbook()
        self.wb.remove(self.wb.active)

    def sheet(self, title: str, heading: str) -> Worksheet:
        ws = self.wb.create_sheet(title)
        ws["A1"] = marking_text(self.label)
        ws["A1"].font = MARK_FONT
        ws["A2"] = heading
        ws["A2"].font = Font(bold=True, size=13)
        ws.oddHeader.center.text = marking_text(self.label)
        ws.oddFooter.center.text = marking_text(self.label)
        return ws

    def header(self, ws: Worksheet, row: int, titles: list[str]) -> None:
        for i, t in enumerate(titles, start=1):
            c = ws.cell(row=row, column=i, value=t)
            c.font, c.fill = HEAD_FONT, HEAD_FILL
            c.alignment = Alignment(wrap_text=True, vertical="center")
            ws.column_dimensions[get_column_letter(i)].width = max(12, min(40, len(t) + 6))

    def source(self, cell: Any, rid: str | None, page: int | None = None) -> None:
        cell.fill = INPUT_FILL
        if not rid:
            return
        rec = self.ledger.maybe(rid)
        where = rec.source() if rec else "record"
        if page and rec and rec.anchor and rec.anchor.page is None:
            where += f" p.{page}"
        cell.comment = Comment(f"source: {rid} ({where})", "Sovereign AI Workbench")

    def bytes(self) -> bytes:
        props = label_properties(self.label)
        self.wb.properties.title = "Sovereign AI Workbench deliverable"
        self.wb.properties.creator = "Sovereign AI Workbench"
        self.wb.properties.keywords = props["keywords"]
        self.wb.properties.category = props["category"]
        self.wb.properties.description = props["comments"]
        buf = io.BytesIO()
        self.wb.save(buf)
        return buf.getvalue()


def calc_sheet(sh: Sheeter, data: dict[str, Any]) -> None:
    calc = _unwrap(data.get("calc") or data)
    ws = sh.sheet("Calculation", str(calc.get("title", "Calculation")))
    ws["A3"] = f"Formula: {calc.get('expression')}"
    sh.header(ws, 5, ["Symbol", "Description", "Value", "Unit", "Value used in formula"])
    consistent = bool((calc.get("excel") or {}).get("raw_units_consistent", True))
    reg = ureg()
    refs: dict[str, str] = {}
    row = 6
    for sym, v in (calc.get("variables") or {}).items():
        ws.cell(row=row, column=1, value=sym)
        ws.cell(row=row, column=2, value=str(v.get("description", "")))
        value_cell = ws.cell(row=row, column=3, value=float(v["value"]))
        sh.source(value_cell, v.get("record"))
        ws.cell(row=row, column=4, value=str(v.get("unit")))
        if consistent:
            ws.cell(row=row, column=5, value=f"=C{row}")
        else:
            factor = float(reg.Quantity(1.0, _unit(str(v["unit"]))).to_base_units().magnitude)
            ws.cell(row=row, column=5, value=f"=C{row}*{factor!r}")
        refs[sym] = f"E{row}"
        row += 1
    row += 1
    result = calc.get("result") or {}
    ws.cell(row=row, column=1, value=str(calc.get("result_name", "result"))).font = Font(bold=True)
    ws.cell(row=row, column=2, value="Result (live formula)")
    formula = to_excel(str(calc.get("rhs") or str(calc.get("expression", "")).split("=", 1)[-1]), refs)
    if not consistent:
        factor = float(reg.Quantity(1.0, _unit(str(result.get("unit")))).to_base_units().magnitude)
        formula = f"({formula})/{factor!r}"
    ws.cell(row=row, column=3, value=f"=ROUND({formula},4)").font = Font(bold=True)
    ws.cell(row=row, column=4, value=str(result.get("unit", "")))
    steps = sh.wb.create_sheet("Steps")
    steps["A1"] = marking_text(sh.label)
    steps["A1"].font = MARK_FONT
    sh.header(steps, 3, ["Step", "Label", "Formula", "Substitution", "Result"])
    for i, st in enumerate(calc.get("steps") or [], start=4):
        for col, key in enumerate(("step", "label", "formula", "substitution", "result"), start=1):
            steps.cell(row=i, column=col, value=str(st.get(key, "")))
    steps.oddHeader.center.text = marking_text(sh.label)
    steps.oddFooter.center.text = marking_text(sh.label)


def offer_comparison(sh: Sheeter, data: dict[str, Any]) -> None:
    comp = _unwrap(data.get("comparison") or data)
    criteria = comp.get("criteria") or []
    offers = comp.get("offers") or []
    if not criteria or not offers:
        raise ToolError("offer comparison needs criteria and offers")
    cws = sh.sheet("Criteria", "Evaluation criteria (from the tender conditions)")
    sh.header(cws, 3, ["Criterion", "Unit", "Better", "Weight", "Limit", "Limit type"])
    crow: dict[str, int] = {}
    for i, c in enumerate(criteria, start=4):
        crow[c["name"]] = i
        cws.cell(row=i, column=1, value=c["name"])
        cws.cell(row=i, column=2, value=c["unit"])
        cws.cell(row=i, column=3, value=c["better"])
        sh.source(cws.cell(row=i, column=4, value=float(c["weight"])), c.get("limit_record"))
        if c.get("limit") is not None:
            sh.source(cws.cell(row=i, column=5, value=float(c["limit"])), c.get("limit_record"))
            cws.cell(row=i, column=6, value=c.get("limit_kind") or ("max" if c["better"] == "lower" else "min"))
    ows = sh.sheet("Offers", "Offer comparison")
    n = len(criteria)
    titles = ["Vendor", "Document"] + [f"{c['name']} ({c['unit']})" for c in criteria] + ["Deviations", "Compliant"]
    titles += [f"{c['name']} score" for c in criteria] + ["Weighted score", "Rank"]
    sh.header(ows, 3, titles)
    first, last = 4, 3 + len(offers)
    val_col = {c["name"]: 3 + i for i, c in enumerate(criteria)}
    dev_col, comp_col = 3 + n, 4 + n
    score_col = {c["name"]: 5 + n + i for i, c in enumerate(criteria)}
    total_col, rank_col = 5 + 2 * n, 6 + 2 * n
    for r, offer in enumerate(offers, start=first):
        ows.cell(row=r, column=1, value=offer["vendor"])
        ows.cell(row=r, column=2, value=offer["document"])
        for c in criteria:
            v = (offer.get("values") or {}).get(c["name"])
            cell = ows.cell(row=r, column=val_col[c["name"]], value=float(v["value"]) if v else None)
            if v:
                sh.source(cell, v.get("record"), v.get("page"))
        devs = offer.get("deviations") or []
        ows.cell(row=r, column=dev_col, value="; ".join(d["text"] for d in devs) or "None")
        conds = []
        for c in criteria:
            if c.get("limit") is None:
                continue
            ref = f"{get_column_letter(val_col[c['name']])}{r}"
            lim = f"Criteria!$E${crow[c['name']]}"
            kind = c.get("limit_kind") or ("max" if c["better"] == "lower" else "min")
            conds.append(f"{ref}<={lim}" if kind == "max" else f"{ref}>={lim}")
        ows.cell(row=r, column=comp_col, value=f'=IF(AND({",".join(conds)}),"Yes","No")' if conds else "Yes")
        for c in criteria:
            col = get_column_letter(val_col[c["name"]])
            rng = f"${col}${first}:${col}${last}"
            formula = f"=MIN({rng})/{col}{r}" if c["better"] == "lower" else f"={col}{r}/MAX({rng})"
            ows.cell(row=r, column=score_col[c["name"]], value=formula).number_format = "0.000"
        parts = [f"{get_column_letter(score_col[c['name']])}{r}*Criteria!$D${crow[c['name']]}" for c in criteria]
        ows.cell(row=r, column=total_col, value="=" + "+".join(parts)).number_format = "0.000"
        cc, tc = get_column_letter(comp_col), get_column_letter(total_col)
        ows.cell(row=r, column=rank_col, value=(
            f'=IF({cc}{r}="Yes",COUNTIFS(${cc}${first}:${cc}${last},"Yes",${tc}${first}:${tc}${last},'
            f'">"&{tc}{r})+1,"-")'))
    rc = get_column_letter(rank_col)
    ows.cell(row=last + 2, column=1, value="Recommended (rank 1)").font = Font(bold=True)
    ows.cell(row=last + 2, column=2, value=f'=IFERROR(INDEX($A${first}:$A${last},MATCH(1,${rc}${first}:${rc}${last},0)),"none")')


def findings_sheet(sh: Sheeter, data: dict[str, Any]) -> None:
    facts = _unwrap(data.get("findings") or data)
    ws = sh.sheet("Findings", f"Thickness readings {facts.get('equipment_tag', {}).get('value', '')}")
    sh.header(ws, 4, ["Location", "Tag", "Measured", "Unit"])
    rows = facts.get("measurements") or []
    for i, m in enumerate(rows, start=5):
        ws.cell(row=i, column=1, value=m.get("location"))
        ws.cell(row=i, column=2, value=m.get("tag"))
        sh.source(ws.cell(row=i, column=3, value=float(m["value"])), m.get("record"))
        ws.cell(row=i, column=4, value=m.get("unit"))
    last = 4 + len(rows)
    ws.cell(row=last + 2, column=1, value="Minimum measured").font = Font(bold=True)
    ws.cell(row=last + 2, column=3, value=f"=MIN(C5:C{last})")
    limit = data.get("limit")
    checks = data.get("checks")
    if limit is None and isinstance(checks, dict):
        for res in checks.get("results") or []:
            if res.get("rule") == "thickness_vs_limit" and res.get("right_value"):
                qs = find_quantities(str(res["right_value"]))
                if qs:
                    limit = {"value": qs[0].magnitude, "record": res.get("right_record")}
    if isinstance(limit, dict) and limit.get("value") is not None:
        ws.cell(row=last + 3, column=1, value="Minimum allowed (SOP)")
        sh.source(ws.cell(row=last + 3, column=3, value=float(limit["value"])), limit.get("record"))
        ws.cell(row=last + 4, column=1, value="Margin")
        ws.cell(row=last + 4, column=3, value=f"=C{last + 2}-C{last + 3}")
        ws.cell(row=last + 5, column=1, value="Status")
        ws.cell(row=last + 5, column=3, value=f'=IF(C{last + 4}>=0,"OK","Below minimum")')


def table_sheet(sh: Sheeter, data: dict[str, Any]) -> None:
    ws = sh.sheet("Table", str(data.get("title", "Table")))
    columns = [str(c) for c in data.get("columns") or []]
    sh.header(ws, 3, columns)
    for r, row in enumerate(data.get("rows") or [], start=4):
        for c, value in enumerate(row, start=1):
            cell = ws.cell(row=r, column=c, value=value.get("value") if isinstance(value, dict) else value)
            if isinstance(value, dict):
                sh.source(cell, value.get("record"))


def make_xlsx(args: dict[str, Any], ctx: ToolContext) -> ToolResult:
    kind = str(args["kind"])
    label = ctx.label()
    sh = Sheeter(ctx.rt.ledger, label)
    data = dict(args["data"])
    {"calc_sheet": calc_sheet, "offer_comparison": offer_comparison, "findings": findings_sheet,
     "table": table_sheet}[kind](sh, data)
    name = str(args.get("name") or f"{kind.replace('_', '-')}.xlsx")
    rec = ctx.rt.files.write_draft(ctx.workspace, name, sh.bytes(), label, ctx.task_id,
                                   overwrite_ok=bool(ctx.meta.get("action_approved", True)))
    return ToolResult(ok=True, summary=f"rendered drafts/{name} ({label.display()})", files=[rec.id],
                      body={"file": rec.relpath, "file_id": rec.id, "label": label.display(), "kind": kind})


SPECS = [ToolSpec(name="make_xlsx", description="Render a spreadsheet with live formulas into drafts/.",
                  input_schema=SCHEMA, handler=make_xlsx, side_effect=True, output_type="file",
                  budget_key="make_xlsx")]
