"""Classification marking helpers shared by the docx, xlsx and pptx renderers."""

from __future__ import annotations

import json
from pathlib import Path

from workbench.core.labels import Label

MARKING_PLACEHOLDER = "{{MARKING}}"


def marking_text(label: Label) -> str:
    return label.marking()


def label_properties(label: Label) -> dict[str, str]:
    return {"keywords": f"classification={label.marking()}", "category": label.display(),
            "comments": json.dumps({"label": label.model_dump(mode="json")}, sort_keys=True)}


def read_embedded_label(path: Path) -> Label | None:
    suffix = path.suffix.lower()
    comments: str | None = None
    if suffix == ".docx":
        from docx import Document

        comments = Document(str(path)).core_properties.comments
    elif suffix == ".pptx":
        from pptx import Presentation

        comments = Presentation(str(path)).core_properties.comments
    elif suffix == ".xlsx":
        from openpyxl import load_workbook

        comments = load_workbook(str(path)).properties.description
    if not comments:
        return None
    try:
        return Label.parse(json.loads(comments)["label"])
    except (ValueError, KeyError):
        return None


def restamp(path: Path, label: Label) -> None:
    """Rewrite the marking in header, footer and properties after an approved downgrade."""
    suffix = path.suffix.lower()
    props = label_properties(label)
    marking = marking_text(label)
    if suffix == ".docx":
        from docx import Document

        doc = Document(str(path))
        for section in doc.sections:
            for part in (section.header, section.footer):
                for p in part.paragraphs:
                    if p.text.strip():
                        for r in p.runs[1:]:
                            r.text = ""
                        if p.runs:
                            p.runs[0].text = marking
        for p in doc.paragraphs:
            if p.style is not None and p.style.name == "WB Marking":
                for r in p.runs[1:]:
                    r.text = ""
                if p.runs:
                    p.runs[0].text = marking
        cp = doc.core_properties
        cp.keywords, cp.category, cp.comments = props["keywords"], props["category"], props["comments"]
        doc.save(str(path))
    elif suffix == ".xlsx":
        from openpyxl import load_workbook

        wb = load_workbook(str(path))
        for ws in wb.worksheets:
            ws.oddHeader.center.text = marking
            ws.oddFooter.center.text = marking
            if ws["A1"].value and str(ws["A1"].value).startswith(("UNCLASSIFIED", "RESTRICTED", "CONFIDENTIAL",
                                                                   "SECRET")):
                ws["A1"].value = marking
        wb.properties.keywords = props["keywords"]
        wb.properties.category = props["category"]
        wb.properties.description = props["comments"]
        wb.save(str(path))
    elif suffix == ".pptx":
        from pptx import Presentation

        prs = Presentation(str(path))
        for slide in prs.slides:
            for shape in slide.shapes:
                if shape.name == "WB Marking" and shape.has_text_frame:
                    shape.text_frame.text = marking
        cp = prs.core_properties
        cp.keywords, cp.category, cp.comments = props["keywords"], props["category"], props["comments"]
        prs.save(str(path))
    elif suffix in {".md", ".txt", ".py", ".csv"}:
        text = path.read_text(encoding="utf-8")
        lines = text.splitlines()
        if lines and lines[0].startswith(("<!-- classification:", "# classification:")):
            prefix = "<!-- classification:" if suffix == ".md" else "# classification:"
            end = " -->" if suffix == ".md" else ""
            lines[0] = f"{prefix} {marking}{end}"
            path.write_text("\n".join(lines) + "\n", encoding="utf-8")
