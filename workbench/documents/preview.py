"""Readable previews of a stored file, so a document can be checked without downloading it.

The preview is plain structured text: paragraphs, headings and tables. It is produced from the
file itself, never from a cached copy, and it is capped so that a large document cannot fill the
browser. Anything without a text form (an image, an unknown binary) reports its own kind and the
interface shows it directly or offers the download instead.
"""

from __future__ import annotations

import csv
import io
from pathlib import Path
from typing import Any

MAX_CHARS = 120_000
MAX_PAGES = 40
MAX_ROWS = 200
IMAGES = {".svg", ".png", ".jpg", ".jpeg", ".webp", ".gif"}
PLAIN = {".md", ".txt", ".py", ".json", ".yaml", ".yml", ".log", ".csv"}


def _cap(blocks: list[dict[str, Any]], used: int) -> tuple[list[dict[str, Any]], bool]:
    out: list[dict[str, Any]] = []
    for block in blocks:
        size = len(block.get("text") or "") + sum(len(", ".join(r)) for r in block.get("rows") or [])
        if used + size > MAX_CHARS:
            return out, True
        used += size
        out.append(block)
    return out, False


def _docx(path: Path) -> list[dict[str, Any]]:
    import docx

    doc = docx.Document(str(path))
    blocks: list[dict[str, Any]] = []
    for para in doc.paragraphs:
        text = para.text.strip()
        if not text:
            continue
        style = (para.style.name or "").lower() if para.style is not None else ""
        kind = "heading" if style.startswith("heading") or style == "title" else "text"
        blocks.append({"kind": kind, "text": text})
    for table in doc.tables:
        rows = [[cell.text.strip() for cell in row.cells] for row in table.rows[:MAX_ROWS]]
        if rows:
            blocks.append({"kind": "table", "header": rows[0], "rows": rows[1:]})
    return blocks


def _xlsx(path: Path) -> list[dict[str, Any]]:
    from openpyxl import load_workbook

    blocks: list[dict[str, Any]] = []
    book = load_workbook(str(path), data_only=True)
    # A cell whose formula has never been calculated has no cached value, so show the formula itself.
    written = load_workbook(str(path), data_only=False)
    for sheet in book.worksheets:
        source = written[sheet.title]
        rows = []
        for index, row in enumerate(sheet.iter_rows(max_row=MAX_ROWS + 1, values_only=True), start=1):
            cells = []
            for column, value in enumerate(row, start=1):
                if value is None:
                    value = source.cell(row=index, column=column).value
                cells.append("" if value is None else str(value))
            rows.append(cells)
        rows = [r for r in rows if any(c.strip() for c in r)]
        if not rows:
            continue
        blocks.append({"kind": "heading", "text": sheet.title})
        # A sheet has no reliable header row: the first line is often a classification banner.
        blocks.append({"kind": "table", "header": [], "rows": rows})
    book.close()
    written.close()
    return blocks


def _csv(path: Path) -> list[dict[str, Any]]:
    rows = list(csv.reader(io.StringIO(path.read_text(encoding="utf-8", errors="replace"))))
    if not rows:
        return []
    body = rows[1:]
    blocks: list[dict[str, Any]] = [{"kind": "table", "header": rows[0], "rows": body[:MAX_ROWS]}]
    if len(body) > MAX_ROWS:
        blocks.append({"kind": "note", "text": f"Showing the first {MAX_ROWS} of {len(body)} rows."})
    return blocks


def _pdf(path: Path, reader: Any) -> list[dict[str, Any]]:
    blocks: list[dict[str, Any]] = []
    for page in reader.read(path)[:MAX_PAGES]:
        blocks.append({"kind": "page", "text": f"Page {page.page}"})
        for para in [p.strip() for p in page.text.split("\n") if p.strip()]:
            blocks.append({"kind": "text", "text": para})
        for table in page.tables:
            blocks.append({"kind": "table", "header": table.header, "rows": table.rows[:MAX_ROWS]})
    return blocks


def build(path: Path, reader: Any) -> dict[str, Any]:
    """``{kind, blocks, truncated}``: ``kind`` is text, image or binary."""
    suffix = path.suffix.lower()
    if suffix in IMAGES:
        return {"kind": "image", "blocks": [], "truncated": False}
    try:
        if suffix == ".docx":
            blocks = _docx(path)
        elif suffix == ".csv":
            blocks = _csv(path)
        elif suffix == ".xlsx":
            blocks = _xlsx(path)
        elif suffix == ".pdf":
            blocks = _pdf(path, reader)
        elif suffix in PLAIN:
            text = path.read_text(encoding="utf-8", errors="replace")
            blocks = [{"kind": "code" if suffix in {".py", ".json", ".yaml", ".yml"} else "text", "text": text}]
        else:
            return {"kind": "binary", "blocks": [], "truncated": False}
    except Exception:  # a preview must never break the page; the download still works
        return {"kind": "binary", "blocks": [], "truncated": False}
    blocks, truncated = _cap(blocks, 0)
    return {"kind": "text", "blocks": blocks, "truncated": truncated}
